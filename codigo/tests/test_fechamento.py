"""Fechamento do mês para o contador (Excel, PDF e ZIP com comprovantes)."""

import io
import zipfile
from datetime import date, timedelta

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import contexto, fechamento, main
from app.arquivos import RepositorioArquivos
from app.banco import BancoBoletos
from app.fila_gravacao import FilaGravacao
from app.repositorio import RepositorioBoletos

SETEMBRO = date(2026, 9, 1)
HOJE = date(2026, 9, 17)


def _boletos():
    return [
        {"id": 1, "beneficiario": "Energisa", "valor": 480.5, "vencimento": date(2026, 9, 10), "status": "Pago",
         "data_pagamento": date(2026, 9, 9), "centro_custo": "Energia", "pago_por": "Maria", "comprovante": "arquivos/x.pdf"},
        {"id": 2, "beneficiario": "Papelaria", "valor": 100, "vencimento": date(2026, 9, 12), "status": "Pendente", "centro_custo": ""},
        {"id": 3, "beneficiario": "Aluguel", "valor": 2000, "vencimento": date(2026, 9, 30), "status": "Pendente", "centro_custo": "Administrativo"},
        {"id": 4, "beneficiario": "Cancelado", "valor": 50, "vencimento": date(2026, 9, 5), "status": "Cancelado"},
        {"id": 5, "beneficiario": "Agosto pago em setembro", "valor": 70, "vencimento": date(2026, 8, 31), "status": "Pago",
         "data_pagamento": date(2026, 9, 1), "centro_custo": "Energia"},
        {"id": -1, "beneficiario": "Aguardando gravação", "valor": 1, "vencimento": date(2026, 9, 2), "status": "Pendente"},
    ]


def test_totais_por_vencimento_e_por_pagamento():
    pix = [{"id": 1, "recebedor": "João Frete", "valor": 150, "vencimento": date(2026, 9, 20), "status": "Pago",
            "data_pagamento": date(2026, 9, 16), "tipo_chave": "celular", "chave": "+5565999990000", "centro_custo": "Obras"}]
    f = fechamento.montar(_boletos(), pix, SETEMBRO, "vencimento", hoje=HOJE)
    t = f["totais"]
    assert (t["total"], t["pago"], t["em_aberto"], t["atrasado"]) == (2730.5, 630.5, 2100.0, 100.0)
    assert (t["quantidade"], t["boletos"], t["pix"], t["cancelados"], t["com_comprovante"], t["sem_comprovante"]) == (4, 3, 1, 1, 1, 1)
    assert [c["centro_custo"] for c in f["por_centro"]] == ["Administrativo", "Energia", "Obras", fechamento.SEM_CENTRO]
    assert f["titulo"] == "setembro de 2026"

    pagamentos = fechamento.montar(_boletos(), pix, SETEMBRO, "pagamento", hoje=HOJE)
    assert [i["nome"] for i in pagamentos["pagos"]] == ["Agosto pago em setembro", "Energisa", "João Frete"]
    assert pagamentos["totais"]["em_aberto"] == 0

    with open_xlsx(fechamento.gerar_xlsx(f)) as wb:
        assert wb.sheetnames == ["Resumo", "Pagos", "Em aberto", "Cancelados"]
        assert wb["Resumo"]["B5"].value == 630.5
    pdf = fechamento.gerar_pdf(f)
    assert pdf.startswith(b"%PDF") and len(pdf) > 1500


class open_xlsx:
    def __init__(self, dados):
        self.wb = load_workbook(io.BytesIO(dados))

    def __enter__(self):
        return self.wb

    def __exit__(self, *_):
        self.wb.close()


def test_zip_para_o_contador_com_comprovantes(tmp_path):
    comprovante = tmp_path / "x.pdf"
    comprovante.write_bytes(b"%PDF comprovante")
    f = fechamento.montar(_boletos(), [], SETEMBRO, "vencimento", hoje=HOJE)
    dados = fechamento.gerar_zip(f, lambda item: comprovante if item["comprovante"] else None)
    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        nomes = z.namelist()
        assert {"fechamento-2026-09.pdf", "fechamento-2026-09.xlsx", "comprovantes/2026-09-09_Energisa_480,50.pdf"} <= set(nomes)
        assert "Agosto" not in " ".join(nomes)
    sem = fechamento.montar(_boletos(), [], SETEMBRO, "pagamento", hoje=HOJE)
    with zipfile.ZipFile(io.BytesIO(fechamento.gerar_zip(sem, lambda item: None))) as z:
        assert "Agosto pago em setembro" in z.read("comprovantes/SEM-COMPROVANTE.txt").decode("utf-8")


def test_api_do_fechamento(tmp_path, monkeypatch):
    planilha = BancoBoletos(tmp_path / "boletos.db", tmp_path / "boletos.xlsx")
    fila = FilaGravacao(planilha, tmp_path / "fila.json")
    monkeypatch.setattr(contexto, "repositorio", RepositorioBoletos(planilha, fila))
    monkeypatch.setattr(contexto, "repositorio_pix", RepositorioBoletos(planilha, fila, "pix"))
    monkeypatch.setattr(contexto, "arquivos", RepositorioArquivos(tmp_path / "arquivos"))
    hoje = date.today()
    planilha.adicionar({"beneficiario": "Internet", "valor": 99.9, "vencimento": hoje.replace(day=1) + timedelta(days=3),
                        "status": "Pago", "data_pagamento": hoje}, usuario="Ana")
    with TestClient(main.app) as c:
        dados = c.get("/api/relatorios/fechamento").json()
        assert dados["mes"] == hoje.strftime("%Y-%m") and dados["totais"]["pago"] == 99.9 and dados["pagos"][0]["pago_por"] == "Ana"
        assert c.get("/api/relatorios/fechamento", params={"mes": "2026-13"}).status_code == 422
        for formato, tipo in (("pdf", "application/pdf"), ("xlsx", "spreadsheet"), ("zip", "application/zip")):
            resposta = c.get("/api/relatorios/fechamento/arquivo", params={"formato": formato})
            assert resposta.status_code == 200 and tipo in resposta.headers["content-type"]
            assert f"fechamento-{hoje:%Y-%m}.{formato}" in resposta.headers["content-disposition"]
