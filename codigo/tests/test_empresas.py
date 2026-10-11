"""Módulo Empresas: cadastro, fechamento por período, baixa automática, cheques e relatório."""

from datetime import date, datetime

import pytest

from app import contexto
from app.empresas import relatorio, servico


@pytest.fixture(autouse=True)
def banco_limpo(tmp_path, monkeypatch):
    """Cada teste usa um banco de empresas só dele."""
    from app.banco import BancoBoletos
    from app.empresas.configuracao import ConfigEmpresas
    from app.empresas.modelos import ABAS_EMPRESAS

    banco = BancoBoletos(tmp_path / "empresas.db", tmp_path / "fechamentos.xlsx", tmp_path / "backups",
                         abas=ABAS_EMPRESAS, centros=False, nome_backup="empresas")
    monkeypatch.setattr(contexto, "empresas", banco)
    monkeypatch.setattr(contexto, "empresas_config", ConfigEmpresas(tmp_path / "empresas.json"))
    yield banco
    banco.parar()


@pytest.fixture
def imapal():
    return servico.salvar_empresa({"nome": "Imapal", "cnpj": "12.345.678/0001-95", "apelidos": "Imapal Construções"},
                                  "Fábio")


def _itens(nf, pedido, data, material="Brita 01", quantidade=12, unitario=158.80):
    return [{"data": data, "material": material, "unidade": "M³", "quantidade": quantidade,
             "valor_unitario": unitario, "numero_nf": nf, "pedido": pedido}]


# --------------------------------------------------------------------------- período e cadastro


def test_a_nota_entra_no_fechamento_pelo_dia_20():
    assert servico.competencia_de(date(2026, 7, 29), 20) == "2026-08"
    assert servico.competencia_de(date(2026, 8, 17), 20) == "2026-08"
    assert servico.competencia_de(date(2026, 8, 21), 20) == "2026-09"
    assert servico.competencia_de(date(2026, 12, 29), 20) == "2027-01"
    assert servico.periodo("2026-08", 20) == (date(2026, 7, 21), date(2026, 8, 20))
    assert servico.periodo("2026-01", 20) == (date(2025, 12, 21), date(2026, 1, 20))


def test_identifica_a_empresa_por_cnpj_nome_ou_apelido(imapal):
    assert servico.identificar_empresa(cnpj="12345678000195")["id"] == imapal["id"]
    assert servico.identificar_empresa(nome="IMAPAL")["id"] == imapal["id"]
    assert servico.identificar_empresa(nome="Imapal Construções Ltda")["id"] == imapal["id"]
    assert servico.identificar_empresa(texto="Cheque da imapal, pedido 7043")["id"] == imapal["id"]
    assert servico.identificar_empresa(nome="Outra Empresa") is None


def test_nao_deixa_cadastrar_duas_empresas_com_o_mesmo_nome(imapal):
    with pytest.raises(servico.ErroEmpresas):
        servico.salvar_empresa({"nome": "imapal"}, "Fábio")


def test_renomear_a_empresa_leva_junto_as_notas(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 10)), "Fábio")
    servico.salvar_empresa({"nome": "Imapal Materiais"}, "Fábio", imapal["id"])
    assert {n["empresa"] for n in contexto.empresas.listar("notas")} == {"Imapal Materiais"}


# --------------------------------------------------------------------------- fechamento


def test_fechamento_junta_as_notas_do_periodo(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 7, 29)), "Fábio")
    servico.lancar_notas("Imapal", _itens("6951", "7139", date(2026, 8, 14), "Pedrisco"), "Fábio")
    servico.lancar_notas("Imapal", _itens("7001", "7200", date(2026, 8, 25)), "Fábio")  # já é do próximo

    f = servico.fechamento("Imapal", "2026-08", hoje=date(2026, 8, 18))
    assert f["inicio"] == date(2026, 7, 21) and f["fim"] == date(2026, 8, 20)
    assert [l["numero_nf"] for l in f["linhas"]] == ["6859", "6951"]
    assert f["totais"]["total"] == 3811.2 and f["totais"]["falta"] == 3811.2
    assert servico.fechamento("Imapal", "2026-09")["totais"]["total"] == 1905.6


def test_valor_total_sai_da_quantidade_vezes_o_unitario(imapal):
    nota = servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 3)), "Fábio")[0]
    assert nota["valor_total"] == 1905.6 and nota["status"] == "Em aberto"


# --------------------------------------------------------------------------- baixa automática


def test_comprovante_com_o_numero_da_nf_da_baixa_so_naquela_nota(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 3)), "Fábio")
    servico.lancar_notas("Imapal", _itens("6869", "7053", date(2026, 8, 5)), "Fábio")

    recebimento = servico.registrar_recebimento(
        {"empresa": "Imapal", "forma": "Pix", "valor": 1905.60, "data": date(2026, 8, 20), "notas": "6869",
         "pagador": "IMAPAL LTDA", "id_transacao": "E12345678202608201200"}, "Fábio")

    notas = {n["numero_nf"]: n for n in contexto.empresas.listar("notas")}
    assert notas["6869"]["status"] == "Recebido" and notas["6869"]["valor_recebido"] == 1905.6
    assert notas["6859"]["status"] == "Em aberto"
    assert recebimento["aplicacoes"].startswith("#2 NF 6869: R$ 1.905,60")
    assert servico.fechamento("Imapal", "2026-08")["totais"]["falta"] == 1905.6


def test_pagamento_parcial_deixa_a_nota_parcial_e_mostra_quanto_falta(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 3)), "Fábio")
    servico.registrar_recebimento({"empresa": "Imapal", "valor": 900, "data": date(2026, 8, 10), "notas": "6859"}, "Ana")

    nota = contexto.empresas.listar("notas")[0]
    assert nota["status"] == "Parcial" and nota["valor_recebido"] == 900
    assert servico.fechamento("Imapal", "2026-08")["totais"] == {
        "quantidade": 1, "total": 1905.6, "recebido": 900.0, "falta": 1005.6, "em_cheque": 0.0, "notas": 1}


def test_sem_numero_de_nf_o_valor_baixa_as_notas_mais_antigas(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 3)), "Fábio")
    servico.lancar_notas("Imapal", _itens("6869", "7053", date(2026, 8, 5)), "Fábio")
    recebimento = servico.registrar_recebimento({"empresa": "Imapal", "valor": 2500, "data": date(2026, 8, 20)}, "Ana")

    notas = {n["numero_nf"]: n for n in contexto.empresas.listar("notas")}
    assert notas["6859"]["status"] == "Recebido"
    assert notas["6869"]["status"] == "Parcial" and notas["6869"]["valor_recebido"] == 594.4
    assert recebimento["sobra"] == 0


def test_sobra_vira_credito_e_cancelar_devolve_tudo(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 3)), "Fábio")
    recebimento = servico.registrar_recebimento({"empresa": "Imapal", "valor": 2000, "data": date(2026, 8, 20)}, "Ana")
    assert recebimento["sobra"] == 94.4 and "sobra (crédito)" in recebimento["aplicacoes"]

    servico.cancelar_recebimento(recebimento["id"], "Ana")
    nota = contexto.empresas.listar("notas")[0]
    assert nota["status"] == "Em aberto" and not nota["valor_recebido"]
    assert contexto.empresas.obter(recebimento["id"], "recebimentos")["status"] == "Cancelado"


def test_recebimento_descobre_a_empresa_pelo_numero_da_nota(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 3)), "Fábio")
    recebimento = servico.registrar_recebimento({"valor": 1905.60, "notas": "6859", "data": date(2026, 8, 21)}, "Ana")
    assert recebimento["empresa"] == "Imapal"


# --------------------------------------------------------------------------- cheques


def test_cheque_prende_as_notas_do_pedido_e_libera_ao_compensar(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 3)), "Fábio")
    cheque = servico.lancar_cheque({"empresa": "Imapal", "numero": "000123", "vencimento": date(2026, 9, 10),
                                    "valor": 1905.60, "banco": "Sicredi", "agencia": "0710", "coop": "3003",
                                    "conta": "12345", "pedidos": "7043"}, "Fábio")
    nota = contexto.empresas.listar("notas")[0]
    assert nota["status"] == "Em cheque" and nota["cheque_id"] == cheque["id"]
    assert servico.fechamento("Imapal", "2026-08")["totais"]["em_cheque"] == 1905.6

    servico.atualizar_cheque(cheque["id"], {"status": "Compensado", "data_deposito": date(2026, 9, 10)}, "Fábio")
    nota = contexto.empresas.listar("notas")[0]
    assert nota["status"] == "Recebido" and nota["valor_recebido"] == 1905.6
    recebimento = contexto.empresas.listar("recebimentos")[0]
    assert recebimento["forma"] == "Cheque" and recebimento["cheque_id"] == cheque["id"]


def test_cheque_devolvido_solta_as_notas(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 8, 3)), "Fábio")
    cheque = servico.lancar_cheque({"empresa": "Imapal", "valor": 1905.60, "vencimento": date(2026, 9, 10),
                                    "pedidos": "7043"}, "Fábio")
    servico.atualizar_cheque(cheque["id"], {"status": "Devolvido"}, "Fábio")
    assert contexto.empresas.listar("notas")[0]["status"] == "Em aberto"


def test_cheques_a_depositar_viram_avisos_de_vencimento(imapal):
    servico.lancar_cheque({"empresa": "Imapal", "numero": "123", "valor": 500, "vencimento": date(2026, 9, 19)}, "Fábio")
    alertas = servico.alertas_de_cheques(hoje=date(2026, 9, 18), dias_antes=3)
    assert alertas[0]["tipo_pagamento"] == "cheque" and alertas[0]["valor"] == 500
    assert "Imapal" in alertas[0]["beneficiario"]
    assert servico.alertas_de_cheques(hoje=date(2026, 9, 1), dias_antes=3) == []


# --------------------------------------------------------------------------- relatório


def test_relatorio_sai_no_formato_da_planilha_usada_hoje(imapal):
    from io import BytesIO

    from openpyxl import load_workbook

    for nf, pedido, dia, material in (("6859", "7043", 29, "Brita 01"), ("6869", "7053", 30, "Pedrisco")):
        servico.lancar_notas("Imapal", _itens(nf, pedido, date(2026, 7, dia), material), "Fábio")
    f = servico.fechamento("Imapal", "2026-08", hoje=date(2026, 8, 18))
    conteudo = relatorio.gerar_xlsx(f, {"nome": "Prospera Industria E Serviços LTDA", "cnpj": "03.880.985/0001-85"})

    ws = load_workbook(BytesIO(conteudo)).active
    assert ws["A1"].value == "Prospera Industria E Serviços LTDA 03.880.985/0001-85"
    assert ws["A2"].value == "Fechamento Imapal"
    assert [c.value for c in ws[5]] == ["DATA", "MATERIAL", "UNIDADE", "QUANT.", "VALOR UNI.", "VALOR TOTAL",
                                        "NOTA FISCAL", "Nº PEDIDO"]
    assert [ws.cell(row=6, column=c).value for c in (1, 2, 3, 4, 6, 7, 8)] == [
        datetime(2026, 7, 29), "Brita 01", "M³", 12, 1905.6, "6859", "7043"]
    assert relatorio.nome_do_arquivo(f, "pdf") == "fechamento-imapal-2026-08.pdf"
    assert len(relatorio.gerar_pdf(f, {"nome": "Prospera", "cnpj": "03.880.985/0001-85"})) > 1000


def test_resumo_das_empresas_mostra_atraso_e_cheques(imapal):
    servico.lancar_notas("Imapal", _itens("6859", "7043", date(2026, 7, 10)), "Fábio")  # fechamento de julho
    servico.lancar_notas("Imapal", _itens("6951", "7139", date(2026, 8, 14)), "Fábio")
    servico.lancar_cheque({"empresa": "Imapal", "valor": 300, "vencimento": date(2026, 9, 5)}, "Fábio")

    resumo = servico.resumo_empresas(hoje=date(2026, 8, 18))[0]
    assert resumo["nome"] == "Imapal" and resumo["competencia"] == "2026-08"
    assert resumo["em_aberto"] == 3811.2 and resumo["atrasado"] == 1905.6
    assert resumo["cheques"] == {"quantidade": 1, "valor": 300.0}
