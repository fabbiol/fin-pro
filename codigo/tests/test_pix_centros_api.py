"""Pix, centros de custo, análise de gastos e endpoints do aplicativo de desktop."""

from datetime import date, timedelta
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import contexto, janela, main
from app.banco import BancoBoletos
from app.fila_gravacao import FilaGravacao
from app.pix import gerar_copia_e_cola
from app.planilha import PlanilhaBoletos
from app.repositorio import FonteLembretes, RepositorioBoletos
from app.servicos import analise_gastos

from gerar_boleto_exemplo import montar_codigo


@pytest.fixture
def cliente(tmp_path, monkeypatch):
    planilha = BancoBoletos(tmp_path / "boletos.db", tmp_path / "boletos.xlsx", tmp_path / "backups")
    fila = FilaGravacao(planilha, tmp_path / "fila.json")
    monkeypatch.setattr(contexto, "planilha", planilha)
    monkeypatch.setattr(contexto, "fila", fila)
    monkeypatch.setattr(contexto, "repositorio", RepositorioBoletos(planilha, fila))
    monkeypatch.setattr(contexto, "repositorio_pix", RepositorioBoletos(planilha, fila, "pix"))
    monkeypatch.setattr(main.preferencias.config, "DADOS_DIR", tmp_path)
    with TestClient(main.app) as c:
        yield c


def test_lanca_edita_e_lista_pix(cliente):
    amanha = (date.today() + timedelta(days=1)).isoformat()
    r = cliente.post("/api/pix", json={"tipo_chave": "celular", "chave": "(65) 99603-5105", "valor": 250.5,
                                      "recebedor": "João Serviços", "vencimento": amanha, "centro_custo": "Manutenção"})
    assert r.status_code == 201, r.text
    pix = r.json()
    assert pix["chave"] == "+5565996035105" and pix["chave_formatada"] == "(65) 99603-5105"
    assert pix["tipo_chave_rotulo"] == "Celular" and pix["situacao"] == "a_vencer"

    assert cliente.post("/api/pix", json={"tipo_chave": "cpf_cnpj", "chave": "111.111.111-11", "valor": 10}).status_code == 422
    assert cliente.post("/api/pix", json={"tipo_chave": "email", "chave": "a@b.com"}).json()["detail"]["codigo"] == "valor_obrigatorio"

    editado = cliente.put(f"/api/pix/{pix['id']}", json={"tipo_chave": "email", "chave": "Financeiro@Empresa.com.br"}).json()
    assert editado["chave"] == "financeiro@empresa.com.br" and editado["valor"] == 250.5
    pago = cliente.patch(f"/api/pix/{pix['id']}/status", json={"status": "Pago"}).json()
    assert pago["data_pagamento"] == date.today().isoformat()

    lista = cliente.get("/api/pix", params={"q": "joao", "centro": "manutenção"}).json()
    assert lista["total"] == 1 and lista["resumo"]["valor_pago"] == 250.5
    copia = cliente.get(f"/api/pix/{pix['id']}/copia-e-cola").json()
    assert copia["gerado"] and copia["copia_e_cola"].startswith("000201")
    assert cliente.get(f"/api/pix/{pix['id']}/qrcode.svg").headers["content-type"].startswith("image/svg")
    assert cliente.get("/api/pix/exportar", params={"formato": "xlsx"}).status_code == 200

    assert contexto.planilha.exportar()  # o Excel é gerado a partir do banco
    wb = load_workbook(contexto.planilha.caminho)
    assert wb["Pix"]["C2"].value == "João Serviços" and wb["Pix"]["D2"].value == "E-mail"
    assert any(l[1] == "Lançamento (Pix)" for l in wb["Histórico"].iter_rows(values_only=True))
    assert cliente.delete(f"/api/pix/{pix['id']}").status_code == 204


def test_pix_copia_e_cola_e_duplicado(cliente):
    codigo = gerar_copia_e_cola("cpf_cnpj", "02.359.273/0001-52", 360, "Hiper JN", txid="NF672362")
    lido = cliente.post("/api/pix/ler-codigo", json={"texto": codigo}).json()
    assert lido["tipo_chave"] == "cpf_cnpj" and lido["valor"] == 360 and lido["duplicado"] is None
    r = cliente.post("/api/pix", json={"copia_e_cola": codigo})
    assert r.status_code == 201 and r.json()["recebedor"] == "HIPER JN" and r.json()["txid"] == "NF672362"
    assert cliente.post("/api/pix", json={"copia_e_cola": codigo}).status_code == 409
    assert cliente.post("/api/pix/ler-codigo", json={"texto": codigo}).json()["duplicado"]["id"] == r.json()["id"]
    alterado = codigo[:-1] + ("0" if codigo[-1] != "0" else "1")
    assert cliente.post("/api/pix", json={"copia_e_cola": alterado}).json()["detail"]["codigo"] == "copia_e_cola_invalido"


def test_centros_de_custo_e_renomear(cliente):
    centros = cliente.get("/api/centros").json()["centros"]
    assert any(c["nome"] == "Combustível e Veículos" for c in centros)

    _, linha = montar_codigo("001", date.today() + timedelta(days=5), Decimal("150.00"), "3" * 25)
    boleto = cliente.post("/api/boletos", json={"linha_digitavel": linha, "beneficiario": "Posto Central", "valor": 150.0,
                                               "centro_custo": "  Frota   Pesada "}).json()
    assert boleto["centro_custo"] == "Frota Pesada"
    centros = {c["nome"]: c for c in cliente.get("/api/centros").json()["centros"]}
    assert centros["Frota Pesada"]["boletos"] == 1 and centros["Frota Pesada"]["total"] == 150.0  # criado automaticamente

    assert cliente.post("/api/centros", json={"nome": "frota pesada"}).status_code == 422  # já existe
    r = cliente.put("/api/centros", json={"nome_anterior": "Frota Pesada", "nome": "Veículos Pesados", "descricao": "Caminhões"})
    assert r.status_code == 200, r.text
    assert cliente.get(f"/api/boletos/{boleto['id']}").json()["centro_custo"] == "Veículos Pesados"
    assert cliente.get("/api/boletos", params={"centro": "veículos pesados"}).json()["total"] == 1
    assert cliente.get("/api/boletos", params={"centro": "__sem__"}).json()["total"] == 0

    cliente.put("/api/centros", json={"nome_anterior": "Outros", "nome": "Outros", "ativo": False})
    assert not {c["nome"]: c for c in cliente.get("/api/centros").json()["centros"]}["Outros"]["ativo"]


def test_analise_de_gastos(cliente):
    hoje = date.today()
    for valor, centro in ((100.0, "Obras"), (300.0, "Obras"), (50.0, "")):
        _, linha = montar_codigo("237", hoje + timedelta(days=int(valor)), Decimal(str(valor)), str(int(valor))[0] * 25)
        cliente.post("/api/boletos", json={"linha_digitavel": linha, "beneficiario": f"Fornecedor {valor}", "valor": valor, "centro_custo": centro})
    cliente.post("/api/pix", json={"tipo_chave": "email", "chave": "a@b.com", "valor": 80, "centro_custo": "Obras",
                                   "vencimento": hoje.isoformat(), "status": "Pago"})
    analise = cliente.get("/api/gastos").json()
    assert analise["total"] == 530.0 and analise["total_pix"] == 80.0 and analise["sem_centro"] == 50.0
    obras = analise["centros"][0]
    assert obras["centro"] == "Obras" and obras["total"] == 480.0 and obras["pix"] == 80.0 and obras["percentual"] == 90.6
    assert cliente.get("/api/gastos", params={"situacao": "pagos"}).json()["total"] == 80.0
    assert cliente.get("/api/gastos", params={"tipo": "boleto"}).json()["total"] == 450.0
    exportado = cliente.get("/api/gastos/exportar")
    assert exportado.status_code == 200 and exportado.content[:2] == b"PK"


def test_analise_por_mes_agrupa_centros_menores():
    hoje = date(2026, 9, 14)
    boletos = [{"id": i, "valor": 10.0 * (i + 1), "vencimento": hoje, "status": "Pendente", "centro_custo": f"C{i}"} for i in range(8)]
    analise = analise_gastos(boletos, [], hoje=hoje)
    assert len(analise["series"]) == 7 and analise["series"][-1] == "Outros centros"
    assert analise["serie_mensal"][0]["centros"]["Outros centros"] == 30.0  # C0 (10) + C1 (20)


def test_notificacoes_incluem_pix(tmp_path):
    planilha = BancoBoletos(tmp_path / "b.db", tmp_path / "b.xlsx")
    fila = FilaGravacao(planilha, tmp_path / "f.json")
    pix = RepositorioBoletos(planilha, fila, "pix")
    pix.adicionar({"recebedor": "Maria", "tipo_chave": "celular", "chave": "+5565996035105", "valor": 90.0,
                   "vencimento": date.today()})
    fonte = FonteLembretes(RepositorioBoletos(planilha, fila), pix)
    [registro] = fonte.listar()
    assert registro["id"] == "pix-1" and registro["chave_pix"] == "(65) 99603-5105"
    from app.notificacoes.regras import montar_alertas, montar_mensagem

    alertas = montar_alertas(fonte.listar(), date.today())
    mensagem = montar_mensagem(alertas, date.today(), "http://localhost:8000")
    assert mensagem.titulo == "Pix vence hoje" and "Chave Pix: (65) 99603-5105" in mensagem.texto


def test_migra_planilha_antiga(tmp_path):
    from openpyxl import Workbook

    caminho = tmp_path / "antiga.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Boletos"
    ws.append(["ID", "Data de Lançamento", "Beneficiário", "Valor (R$)", "Status"])
    ws.append([1, None, "Antigo", 10.0, "Pendente"])
    wb.create_sheet("Histórico")
    wb.save(caminho)
    planilha = PlanilhaBoletos(caminho)
    assert planilha.migrar()
    nomes = load_workbook(caminho).sheetnames
    assert nomes == ["Boletos", "Pix", "Centros de Custo", "Histórico"]
    assert not planilha.migrar()
    assert planilha.listar()[0]["centro_custo"] == "" and planilha.listar("pix") == []


def test_endpoints_do_aplicativo(cliente, monkeypatch):
    estado = cliente.get("/api/app").json()
    assert estado["modo_app"] is False and estado["preferencias"]["fechar_para_bandeja"] is True
    assert cliente.post("/api/app/mostrar", json={"rota": "pix"}).json() == {"modo_app": False}
    assert cliente.get("/api/app/abrir", params={"rota": "boletos?pagar=3"}, follow_redirects=False).headers["location"] == "/#/boletos?pagar=3"

    rotas = []
    janela.registrar(rotas.append)
    try:
        assert cliente.post("/api/app/mostrar", json={"rota": "pix?pagar=2"}).json() == {"modo_app": True}
        assert cliente.post("/api/app/mostrar", json={"rota": "javascript:alert(1)"}).json() == {"modo_app": True}
        assert "foi aberto" in cliente.get("/api/app/abrir", params={"rota": "gastos"}).text
        assert rotas == ["pix?pagar=2", None, "gastos"]
        assert main._link_notificacao("pix?pagar=2").endswith("/api/app/abrir?rota=pix%3Fpagar%3D2")
    finally:
        janela.desregistrar()

    prefs = cliente.put("/api/app/preferencias", json={"fechar_para_bandeja": False}).json()["preferencias"]
    assert prefs["fechar_para_bandeja"] is False


def test_painel_mostra_boletos_e_pix_vencidos(cliente):
    hoje = date.today()
    for dias, nome in ((-20, "Pix antigo"), (-3, "Pix recente"), (5, "Pix futuro")):
        r = cliente.post("/api/pix", json={"tipo_chave": "email", "chave": "financeiro@fornecedor.com", "valor": 100 + dias,
                                          "recebedor": nome, "vencimento": (hoje + timedelta(days=dias)).isoformat()})
        assert r.status_code == 201, r.text
    for dias in range(-12, -4):  # mais de 6 boletos vencidos: todos aparecem
        cliente.post("/api/boletos", json={"beneficiario": f"Fornecedor {dias}", "valor": 50, "ignorar_validacao": True,
                                           "vencimento": (hoje + timedelta(days=dias)).isoformat()})
    painel = cliente.get("/api/painel").json()
    k = painel["indicadores"]
    assert (k["vencidos_qtd"], k["vencidos_pix_qtd"], k["vencidos_boletos_qtd"]) == (10, 2, 8)
    vencidos = painel["vencidos"]
    assert len(vencidos) == 10 and vencidos[0]["beneficiario"] == "Pix antigo" and vencidos[0]["tipo_pagamento"] == "pix"
    assert {v["beneficiario"] for v in vencidos if v["tipo_pagamento"] == "pix"} == {"Pix antigo", "Pix recente"}
    assert any(p["beneficiario"] == "Pix futuro" for p in painel["proximos"])


def test_previsao_de_saidas_por_semana(cliente):
    from app.servicos import previsao_saidas

    hoje = date(2026, 9, 17)
    boletos = [{"id": 1, "status": "Pendente", "valor": 100, "vencimento": date(2026, 9, 10), "beneficiario": "Atrasado"},
               {"id": 2, "status": "Pendente", "valor": 200, "vencimento": date(2026, 9, 17), "beneficiario": "Hoje"},
               {"id": 3, "status": "Pendente", "valor": 300, "vencimento": date(2026, 9, 24), "beneficiario": "Semana 2"},
               {"id": 4, "status": "Pago", "valor": 999, "vencimento": date(2026, 9, 18)},
               {"id": 5, "status": "Pendente", "valor": 50, "vencimento": None}]
    pix = [{"id": 1, "status": "Pendente", "valor": 25.5, "vencimento": date(2026, 9, 23), "recebedor": "Frete"}]
    p = previsao_saidas(boletos, pix, hoje, semanas=4)
    assert [s["total"] for s in p["semanas"]] == [225.5, 300.0, 0.0, 0.0]
    assert p["semanas"][0]["pix"] == 25.5 and [i["nome"] for i in p["semanas"][0]["itens"]] == ["Hoje", "Frete"]
    assert (p["em_atraso"]["total"], p["sem_vencimento"]["quantidade"]) == (100.0, 1)
    assert p["totais"] == {"7_dias": 225.5, "30_dias": 525.5, "periodo": 525.5, "com_atraso": 625.5}
    assert cliente.get("/api/previsao", params={"semanas": 12}).json()["semanas"].__len__() == 12
