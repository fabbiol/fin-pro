"""Relatório geral do mês (menu Relatórios): pagos, juros e multas, últimos 12 meses e as partes de cada módulo."""

import io
from datetime import date, timedelta

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import contexto, main, relatorio_geral
from app.arquivos import RepositorioArquivos
from app.banco import BancoBoletos
from app.fila_gravacao import FilaGravacao
from app.repositorio import RepositorioBoletos

SETEMBRO = date(2026, 9, 1)
HOJE = date(2026, 9, 24)


class _Repo:
    def __init__(self, registros):
        self.registros = registros

    def listar(self):
        return list(self.registros)


def _boleto(id_, valor, juros=0.0, pago=date(2026, 9, 21), venc=date(2026, 9, 10), status="Pago"):
    return {"id": id_, "beneficiario": f"Fornecedor {id_}", "valor": valor, "status": status, "vencimento": venc,
            "data_pagamento": pago if status == "Pago" else None, "valor_pago": valor + juros if status == "Pago" else None,
            "juros_multa": juros, "centro_custo": "Manutenção"}


def _preparar(monkeypatch):
    boletos = [_boleto(1, 900, 56.70), _boleto(2, 210), _boleto(3, 500, status="Pendente", venc=date(2026, 9, 30)),
               _boleto(4, 300, 9.90, pago=date(2026, 8, 12), venc=date(2026, 8, 5))]
    pix = [{"id": 1, "recebedor": "João Frete", "valor": 150, "vencimento": date(2026, 9, 18), "status": "Pago",
            "data_pagamento": date(2026, 9, 18), "valor_pago": 150, "centro_custo": "Fretes"}]
    monkeypatch.setattr(contexto, "repositorio", _Repo(boletos))
    monkeypatch.setattr(contexto, "repositorio_pix", _Repo(pix))
    monkeypatch.setattr("app.caixa.servico.saidas_para_gastos", lambda: [])


def test_pagos_juros_e_historico_do_mes(monkeypatch):
    _preparar(monkeypatch)
    r = relatorio_geral.montar(SETEMBRO, lambda p: p == "gastos", HOJE)
    c = r["contas"]
    assert (c["pago"], c["quantidade_pagos"], c["juros_multa"]) == (1316.7, 3, 56.7)
    assert c["por_tipo"]["Boleto"] == {"valor": 1166.7, "quantidade": 2} and c["por_tipo"]["Pix"]["valor"] == 150
    assert [j["nome"] for j in c["juros"]] == ["Fornecedor 1"]
    assert c["vencimento"]["em_aberto"] == 500
    assert len(r["historico"]) == 12 and r["historico"][-1]["mes"] == "2026-09"
    assert r["historico"][-2]["juros_multa"] == 9.9
    assert (r["titulo"], r["anterior"], r["seguinte"]) == ("setembro de 2026", "2026-08", "2026-10")
    assert "caixa" not in r and "diesel" not in r and "empresas" not in r  # sem a permissão, a parte não aparece


def test_sem_permissao_de_gastos_nao_mostra_os_pagamentos(monkeypatch):
    _preparar(monkeypatch)
    r = relatorio_geral.montar(SETEMBRO, lambda p: False, HOJE)
    assert "contas" not in r and "historico" not in r and "previsao" in r


def test_uma_parte_com_erro_nao_derruba_o_relatorio(monkeypatch):
    _preparar(monkeypatch)

    def quebra(*_):
        raise RuntimeError("banco do diesel indisponível")

    monkeypatch.setattr(relatorio_geral, "_diesel", quebra)
    r = relatorio_geral.montar(SETEMBRO, lambda p: p in ("gastos", "diesel"), HOJE)
    assert r["diesel"] == {"erro": "banco do diesel indisponível"} and r["contas"]["pago"] == 1316.7


def test_excel_tem_resumo_juros_e_12_meses(monkeypatch):
    _preparar(monkeypatch)
    r = relatorio_geral.montar(SETEMBRO, lambda p: p == "gastos", HOJE)
    wb = load_workbook(io.BytesIO(relatorio_geral.gerar_xlsx(r)))
    assert wb.sheetnames == ["Resumo", "Juros e multas", "Últimos 12 meses"]
    assert wb["Resumo"]["B5"].value == 1316.7
    assert wb["Juros e multas"]["B5"].value == "Fornecedor 1" and wb["Juros e multas"]["G5"].value == 56.7
    assert wb["Últimos 12 meses"]["A16"].value == "setembro de 2026"


def test_api_do_relatorio_geral(tmp_path, monkeypatch):
    planilha = BancoBoletos(tmp_path / "boletos.db", tmp_path / "boletos.xlsx")
    fila = FilaGravacao(planilha, tmp_path / "fila.json")
    monkeypatch.setattr(contexto, "repositorio", RepositorioBoletos(planilha, fila))
    monkeypatch.setattr(contexto, "repositorio_pix", RepositorioBoletos(planilha, fila, "pix"))
    monkeypatch.setattr(contexto, "arquivos", RepositorioArquivos(tmp_path / "arquivos"))
    hoje = date.today()
    planilha.adicionar({"beneficiario": "Internet", "valor": 99.9, "vencimento": hoje.replace(day=1) + timedelta(days=3),
                        "status": "Pago", "data_pagamento": hoje, "valor_pago": 104.9, "juros_multa": 5}, usuario="Ana")
    with TestClient(main.app) as c:
        dados = c.get("/api/relatorio-geral").json()
        assert dados["mes"] == hoje.strftime("%Y-%m")
        assert dados["contas"]["pago"] == 104.9 and dados["contas"]["juros_multa"] == 5
        for parte in ("caixa", "diesel", "empresas"):
            assert parte in dados and "erro" not in dados[parte], dados.get(parte)
        assert c.get("/api/relatorio-geral", params={"mes": "2026-13"}).status_code == 422
        resposta = c.get("/api/relatorio-geral/arquivo")
        assert resposta.status_code == 200 and "spreadsheet" in resposta.headers["content-type"]
        assert f"relatorio-geral-{hoje:%Y-%m}.xlsx" in resposta.headers["content-disposition"]
