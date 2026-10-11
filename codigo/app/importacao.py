"""Importação de boletos exportados de outros sistemas (ERP, financeiro, contas a pagar).

Aceita Excel (.xlsx/.xlsm e .xls — inclusive o ".xls" em HTML que muitos sistemas geram), CSV e TXT.
As colunas são reconhecidas pelo nome, com os sinônimos mais comuns ("Fornecedor", "Vencto", "Valor Título"...),
e o usuário confere o mapeamento antes de importar. Cada linha é validada (linha digitável, valor, vencimento)
e comparada com a planilha para não lançar o mesmo boleto duas vezes. O mapeamento escolhido fica salvo e é
reaproveitado na próxima importação de um arquivo com as mesmas colunas.
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import threading
import time
import uuid
from collections import OrderedDict
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .boleto.codigo import analisar, somente_digitos
from .boleto.texto import normalizar

MAX_LINHAS = 5000
MAX_COLUNAS = 80
LINHAS_BUSCA_CABECALHO = 30
HORAS_GUARDAR_ARQUIVOS = 24

CAMPOS = [
    {"chave": "linha_digitavel", "rotulo": "Linha digitável / código de barras"},
    {"chave": "beneficiario", "rotulo": "Beneficiário / fornecedor"},
    {"chave": "documento_beneficiario", "rotulo": "CPF/CNPJ do beneficiário"},
    {"chave": "numero_documento", "rotulo": "Nº do documento"},
    {"chave": "vencimento", "rotulo": "Vencimento"},
    {"chave": "valor", "rotulo": "Valor"},
    {"chave": "status", "rotulo": "Situação (pago / em aberto)"},
    {"chave": "data_pagamento", "rotulo": "Data de pagamento"},
    {"chave": "centro_custo", "rotulo": "Centro de custo"},
    {"chave": "observacoes", "rotulo": "Observações"},
]
CHAVES_CAMPOS = [c["chave"] for c in CAMPOS]

# Nomes de coluna usados pelos sistemas (já normalizados: minúsculas, sem acentos e sem pontuação)
APELIDOS: dict[str, list[str]] = {
    "linha_digitavel": ["linha digitavel", "codigo de barras", "codigo barras", "cod barras", "cod de barras", "barras",
                        "linha", "codigo do boleto", "representacao numerica", "ld", "barcode", "linha digitavel boleto"],
    "beneficiario": ["beneficiario", "nome do beneficiario", "cedente", "fornecedor", "nome do fornecedor", "nome fornecedor",
                     "favorecido", "credor", "razao social", "emitente", "empresa", "sacador avalista", "nome"],
    "documento_beneficiario": ["cpf cnpj", "cnpj cpf", "cnpj", "cpf", "cnpj do beneficiario", "cnpj fornecedor",
                               "cnpj do fornecedor", "documento do beneficiario", "cpf cnpj fornecedor"],
    "numero_documento": ["numero do documento", "numero documento", "n do documento", "no do documento", "n documento",
                         "no documento", "num documento", "nro documento", "documento", "doc", "nota fiscal", "nf",
                         "numero nf", "titulo", "numero do titulo", "duplicata", "fatura"],
    "vencimento": ["vencimento", "data de vencimento", "data vencimento", "dt vencimento", "dt venc", "data venc", "venc",
                   "vencto", "dt vencto", "vencimento original"],
    "valor": ["valor", "valor do documento", "valor documento", "valor do titulo", "valor titulo", "valor original",
              "valor a pagar", "valor boleto", "valor do boleto", "valor liquido", "valor total", "vlr", "vlr documento",
              "vl documento", "vlr titulo", "vl titulo", "total"],
    "status": ["status", "situacao", "situacao do titulo", "estado", "pago", "quitado"],
    "data_pagamento": ["data de pagamento", "data pagamento", "dt pagamento", "dt pagto", "data pagto", "pagamento",
                       "data da baixa", "data baixa", "dt baixa", "baixa", "pago em", "data liquidacao", "liquidacao"],
    "centro_custo": ["centro de custo", "centro custo", "c custo", "cc", "departamento", "setor", "obra", "projeto",
                     "unidade", "filial"],
    "observacoes": ["observacoes", "observacao", "obs", "historico", "descricao", "complemento", "memo"],
}
# Palavras que descartam uma coluna para o campo (ex.: "Valor Pago" não é o valor do boleto)
EXCLUIR: dict[str, set[str]] = {
    "linha_digitavel": {"nosso"},
    "beneficiario": {"sacado", "pagador", "cpf", "cnpj", "banco", "conta"},
    "documento_beneficiario": {"numero", "nosso"},
    "numero_documento": {"cpf", "cnpj", "nosso", "banco"},
    "vencimento": {"pagamento", "pagto"},
    "valor": {"juros", "multa", "desconto", "abatimento", "acrescimo", "acrescimos", "pago", "mora", "encargos"},
    "status": {"valor", "data", "dt"},
    "data_pagamento": {"valor", "vencimento", "forma", "banco", "conta", "vlr"},
}
MESES = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6, "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}
FORMATOS = {"xlsx": "Excel", "xls": "Excel 97-2003", "html": "Excel (HTML)", "csv": "CSV / texto"}


class ErroImportacao(ValueError):
    pass


# --------------------------------------------------------------------------- conversões


def _vazio(valor) -> bool:
    return valor is None or (isinstance(valor, str) and not valor.strip())


def texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y") if valor.time() == datetime.min.time() else valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, float):
        if valor.is_integer() and abs(valor) < 1e15:
            return str(int(valor))
        return f"{valor:.2f}".replace(".", ",")
    return re.sub(r"\s+", " ", str(valor)).strip()


def normalizar_titulo(valor) -> str:
    t = normalizar(str(valor or ""))
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


RE_DATA_ISO = re.compile(r"^\s*(\d{4})-(\d{1,2})-(\d{1,2})")
RE_DATA_BR = re.compile(r"^\s*(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})(?:[\sT]|$)")
RE_DATA_MES = re.compile(r"^\s*(\d{1,2})[/\-\s]+(?:de\s+)?([a-z]{3})[a-z]*[/\-\s]+(?:de\s+)?(\d{2,4})")


def ler_data(valor) -> date | None:
    if _vazio(valor) or isinstance(valor, bool):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)):
        # Data como número serial do Excel (ex.: 46280)
        return date(1899, 12, 30) + timedelta(days=int(valor)) if 20000 <= valor <= 80000 else None
    t = str(valor).strip()
    try:
        m = RE_DATA_ISO.match(t)
        if m:
            return date(int(m[1]), int(m[2]), int(m[3]))
        m = RE_DATA_BR.match(t)
        if m:
            ano = int(m[3])
            return date(ano + 2000 if ano < 100 else ano, int(m[2]), int(m[1]))
        m = RE_DATA_MES.match(normalizar(t))
        if m and m[2] in MESES:
            ano = int(m[3])
            return date(ano + 2000 if ano < 100 else ano, MESES[m[2]], int(m[1]))
        if re.fullmatch(r"\d{8}", t):
            if t[:2] in ("19", "20"):
                return date(int(t[:4]), int(t[4:6]), int(t[6:]))
            return date(int(t[4:]), int(t[2:4]), int(t[:2]))
    except ValueError:
        return None
    return None


def ler_valor(valor) -> float | None:
    if _vazio(valor) or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float, Decimal)):
        return round(abs(float(valor)), 2)
    t = re.sub(r"[^\d,.]", "", str(valor))
    if not re.search(r"\d", t):
        return None
    if "," in t and "." in t:
        # O último separador é o decimal: 1.234,56 (Brasil) ou 1,234.56 (EUA)
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        partes = t.split(",")
        t = "".join(partes[:-1]) + "." + partes[-1]
    elif t.count(".") > 1 or re.fullmatch(r"\d{1,3}\.\d{3}", t):
        t = t.replace(".", "")  # 1.500 = mil e quinhentos
    try:
        return round(abs(float(Decimal(t))), 2)
    except (InvalidOperation, ValueError):
        return None


def ler_status(valor, padrao: str = "Pendente") -> str:
    t = normalizar_titulo(texto(valor))
    if not t:
        return padrao
    if re.search(r"cancel|estorn|exclu|anulad", t):
        return "Cancelado"
    if re.search(r"\bnao\b|aberto|a pagar|pendente|a vencer|vencid|agendad|atrasad", t):
        return "Pendente"
    if re.search(r"pag|liquid|quit|baixad|compensad|efetivad|realizad", t) or t in ("s", "sim", "x", "ok", "true"):
        return "Pago"
    return padrao


def formatar_documento(valor: str) -> str:
    d = somente_digitos(valor)
    if len(d) == 14:
        return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"
    if len(d) == 11:
        return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"
    return valor.strip()


def _brl(valor: float) -> str:
    return "R$ " + f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


# --------------------------------------------------------------------------- leitura dos arquivos


def _com_dados(valores) -> bool:
    return any(not _vazio(v) for v in valores)


def _ler_xlsx(dados: bytes) -> OrderedDict:
    try:
        wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
    except Exception as exc:
        raise ErroImportacao("Não foi possível abrir a planilha Excel. Verifique se o arquivo não está corrompido "
                             "ou protegido por senha.") from exc
    abas = OrderedDict()
    try:
        for ws in wb.worksheets:
            linhas = []
            for numero, valores in enumerate(ws.iter_rows(values_only=True), start=1):
                if numero > MAX_LINHAS + LINHAS_BUSCA_CABECALHO:
                    break
                valores = [v.strip() if isinstance(v, str) else v for v in list(valores)[:MAX_COLUNAS]]
                if _com_dados(valores):
                    linhas.append((numero, valores))
            abas[ws.title] = linhas
    finally:
        wb.close()
    return abas


def _ler_xls(dados: bytes) -> OrderedDict:
    try:
        import xlrd
    except ImportError as exc:  # pragma: no cover
        raise ErroImportacao("Para ler arquivos .xls antigos, abra-o no Excel e salve como .xlsx.") from exc
    try:
        livro = xlrd.open_workbook(file_contents=dados)
    except Exception as exc:
        raise ErroImportacao("Não foi possível abrir o arquivo .xls. Abra-o no Excel e salve como .xlsx.") from exc
    abas = OrderedDict()
    for folha in livro.sheets():
        linhas = []
        for r in range(min(folha.nrows, MAX_LINHAS + LINHAS_BUSCA_CABECALHO)):
            valores = []
            for c in range(min(folha.ncols, MAX_COLUNAS)):
                celula = folha.cell(r, c)
                if celula.ctype == xlrd.XL_CELL_DATE:
                    try:
                        valor = xlrd.xldate_as_datetime(celula.value, livro.datemode)
                    except Exception:
                        valor = celula.value
                elif celula.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                    valor = None
                elif celula.ctype == xlrd.XL_CELL_BOOLEAN:
                    valor = bool(celula.value)
                else:
                    valor = celula.value.strip() if isinstance(celula.value, str) else celula.value
                valores.append(valor)
            if _com_dados(valores):
                linhas.append((r + 1, valores))
        abas[folha.name] = linhas
    return abas


class _LeitorTabelasHtml(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.tabelas: list[list[list[str]]] = []
        self._pilha: list[list[list[str]]] = []
        self._linha: list[str] | None = None
        self._celula: list[str] | None = None
        self._colspan = 1

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._pilha.append([])
        elif tag == "tr" and self._pilha:
            self._linha = []
        elif tag in ("td", "th") and self._linha is not None:
            self._celula = []
            try:
                self._colspan = max(1, min(int(dict(attrs).get("colspan") or 1), 50))
            except ValueError:
                self._colspan = 1
        elif tag == "br" and self._celula is not None:
            self._celula.append(" ")

    def handle_data(self, data):
        if self._celula is not None:
            self._celula.append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._celula is not None and self._linha is not None:
            self._linha.append(re.sub(r"\s+", " ", "".join(self._celula)).strip())
            self._linha.extend([""] * (self._colspan - 1))
            self._celula = None
        elif tag == "tr" and self._linha is not None and self._pilha:
            self._pilha[-1].append(self._linha)
            self._linha = None
        elif tag == "table" and self._pilha:
            self.tabelas.append(self._pilha.pop())


def _decodificar(dados: bytes) -> str:
    if dados[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return dados.decode("utf-16")
    try:
        return dados.decode("utf-8-sig")
    except UnicodeDecodeError:
        return dados.decode("cp1252", errors="replace")


def _ler_html(dados: bytes) -> OrderedDict:
    leitor = _LeitorTabelasHtml()
    leitor.feed(_decodificar(dados))
    leitor.close()
    abas = OrderedDict()
    for tabela in leitor.tabelas:
        linhas = [(n, [c or None for c in linha[:MAX_COLUNAS]]) for n, linha in enumerate(tabela, start=1) if _com_dados(linha)]
        if len(linhas) >= 2:
            abas[f"Tabela {len(abas) + 1}"] = linhas[: MAX_LINHAS + LINHAS_BUSCA_CABECALHO]
    return abas


def _ler_texto(dados: bytes) -> OrderedDict:
    conteudo = _decodificar(dados)
    primeiras = [l for l in conteudo.splitlines() if l.strip()][:20]
    separador = ","
    for candidato in ("\t", ";", "|", ","):
        if primeiras and sum(candidato in l for l in primeiras) >= 0.6 * len(primeiras):
            separador = candidato
            break
    linhas = []
    for numero, valores in enumerate(csv.reader(io.StringIO(conteudo), delimiter=separador), start=1):
        if numero > MAX_LINHAS + LINHAS_BUSCA_CABECALHO:
            break
        valores = [v.strip() or None for v in valores[:MAX_COLUNAS]]
        if _com_dados(valores):
            linhas.append((numero, valores))
    return OrderedDict([("Arquivo", linhas)])


def ler_planilhas(dados: bytes, nome: str = "") -> tuple[str, OrderedDict]:
    """Lê o arquivo e devolve (formato, {aba: [(número da linha, [valores])]})."""
    if dados[:5] == b"%PDF-":
        raise ErroImportacao("Este arquivo é um PDF. Para boletos em PDF ou foto, use a tela Ler boletos.")
    if dados[:3] == b"\xff\xd8\xff" or dados[:8] == b"\x89PNG\r\n\x1a\n":
        raise ErroImportacao("Este arquivo é uma imagem. Para fotos de boletos, use a tela Ler boletos.")
    if dados[:4] == b"PK\x03\x04":
        formato, abas = "xlsx", _ler_xlsx(dados)
    elif dados[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        formato, abas = "xls", _ler_xls(dados)
    elif b"<table" in dados[:400000].lower():
        formato, abas = "html", _ler_html(dados)
    else:
        formato, abas = "csv", _ler_texto(dados)
    abas = OrderedDict((n, l) for n, l in abas.items() if l)
    if not abas:
        raise ErroImportacao("Nenhuma linha com dados foi encontrada no arquivo.")
    return formato, abas


# --------------------------------------------------------------------------- colunas


def _pontuar(campo: str, titulo_norm: str) -> int:
    if not titulo_norm:
        return 0
    palavras = set(titulo_norm.split())
    if palavras & EXCLUIR.get(campo, set()):
        return 0
    melhor = 0
    for apelido in APELIDOS[campo]:
        if titulo_norm == apelido:
            melhor = max(melhor, 100 + len(apelido))
        elif all(p in palavras for p in apelido.split()):
            melhor = max(melhor, 60 + len(apelido))
        elif len(apelido) >= 6 and apelido in titulo_norm:
            melhor = max(melhor, 40 + len(apelido))
    return melhor


def _reconhecidos(valores) -> int:
    return sum(1 for v in valores if isinstance(v, str) and any(_pontuar(c, normalizar_titulo(v)) >= 40 for c in CHAVES_CAMPOS))


def _parece_dado(valor) -> bool:
    if _vazio(valor):
        return False
    if not isinstance(valor, str):
        return True
    return bool(ler_data(valor) or len(somente_digitos(valor)) >= 44 or re.fullmatch(r"[R$\s]*[\d.,]+", valor))


def detectar_cabecalho(linhas) -> tuple[int, int]:
    """(número da linha dos títulos ou 0 se não houver, quantidade de títulos reconhecidos)."""
    melhor, melhor_pontos = 0, 0
    for numero, valores in linhas[:LINHAS_BUSCA_CABECALHO]:
        if any(len(somente_digitos(texto(v))) in (44, 47, 48) for v in valores):
            continue
        reconhecidos = _reconhecidos(valores)
        if reconhecidos >= 2 and reconhecidos > melhor_pontos:
            melhor, melhor_pontos = numero, reconhecidos
    if not melhor and linhas:
        numero, valores = linhas[0]
        preenchidos = [v for v in valores if not _vazio(v)]
        if len(preenchidos) >= 2 and not any(_parece_dado(v) for v in preenchidos):
            return numero, 0
    return melhor, melhor_pontos


def montar_tabela(abas: OrderedDict, aba: str | None = None, linha_cabecalho: int | None = None) -> dict:
    if aba not in abas:
        aba = max(abas, key=lambda nome: (detectar_cabecalho(abas[nome])[1], len(abas[nome])))
    linhas = abas[aba]
    if linha_cabecalho is None:
        linha_cabecalho = detectar_cabecalho(linhas)[0]
    cabecalho = next((v for n, v in linhas if n == linha_cabecalho), None) if linha_cabecalho else None
    if cabecalho is None:
        linha_cabecalho = 0
    dados = [(n, v) for n, v in linhas if n > linha_cabecalho]
    largura = min(max([len(v) for _, v in linhas] + [1]), MAX_COLUNAS)
    titulos, vistos = [], {}
    for i in range(largura):
        titulo = texto(cabecalho[i]) if cabecalho and i < len(cabecalho) else ""
        titulo = titulo or f"Coluna {get_column_letter(i + 1)}"
        vistos[titulo] = vistos.get(titulo, 0) + 1
        titulos.append(titulo if vistos[titulo] == 1 else f"{titulo} ({vistos[titulo]})")
    return {
        "abas": list(abas),
        "aba": aba,
        "linha_cabecalho": linha_cabecalho,
        "colunas": titulos,
        "linhas": dados[:MAX_LINHAS],
        "truncado": len(dados) > MAX_LINHAS,
        "opcoes_cabecalho": [(n, " | ".join(t for t in (texto(x) for x in v) if t)[:160]) for n, v in linhas[:LINHAS_BUSCA_CABECALHO]],
    }


def descrever_colunas(tabela: dict) -> list[dict]:
    colunas = []
    for i, titulo in enumerate(tabela["colunas"]):
        exemplos = []
        for _, valores in tabela["linhas"]:
            if i < len(valores) and not _vazio(valores[i]):
                exemplos.append(texto(valores[i])[:60])
                if len(exemplos) == 3:
                    break
        colunas.append({"indice": i, "titulo": titulo, "exemplos": exemplos})
    return colunas


def sugerir_mapeamento(colunas: list[str], linhas) -> dict[str, int | None]:
    pares = sorted(
        ((pontos, campo, i) for campo in CHAVES_CAMPOS for i, titulo in enumerate(colunas)
         if (pontos := _pontuar(campo, normalizar_titulo(titulo))) >= 40),
        key=lambda p: (-p[0], CHAVES_CAMPOS.index(p[1]), p[2]),
    )
    mapa: dict[str, int] = {}
    usadas: set[int] = set()
    for _, campo, i in pares:
        if campo not in mapa and i not in usadas:
            mapa[campo] = i
            usadas.add(i)
    if "linha_digitavel" not in mapa:
        # Coluna sem título reconhecível, mas cheia de códigos de 44/47/48 dígitos
        for i in range(len(colunas)):
            valores = [texto(v[i]) for _, v in linhas[:200] if i < len(v) and not _vazio(v[i])]
            if i not in usadas and valores and sum(len(somente_digitos(t)) in (44, 47, 48) for t in valores) >= 0.6 * len(valores):
                mapa["linha_digitavel"] = i
                break
    return {campo: mapa.get(campo) for campo in CHAVES_CAMPOS}


def validar_mapeamento(mapeamento: dict, largura: int) -> dict[str, int | None]:
    saida = {}
    for campo in CHAVES_CAMPOS:
        valor = mapeamento.get(campo)
        saida[campo] = valor if isinstance(valor, int) and not isinstance(valor, bool) and 0 <= valor < largura else None
    return saida


# --------------------------------------------------------------------------- validação das linhas


def _chave_sem_codigo(beneficiario, numero_documento, valor, vencimento) -> str | None:
    if not (beneficiario and valor and isinstance(vencimento, date)):
        return None
    documento = somente_digitos(numero_documento) or normalizar_titulo(numero_documento)
    return f"{normalizar_titulo(beneficiario)}|{documento}|{float(valor):.2f}|{vencimento.isoformat()}"


def preparar(tabela: dict, mapeamento: dict, *, existentes: list[dict], centro_padrao: str = "",
             status_padrao: str = "Pendente", nome_arquivo: str = "", hoje: date | None = None) -> list[dict]:
    """Converte as linhas do arquivo em boletos, com estado: pronto, aviso, erro, duplicado ou repetido."""
    hoje = hoje or date.today()
    codigos_planilha, chaves_planilha = {}, {}
    for r in existentes:
        if r.get("status") == "Cancelado":
            continue
        if r.get("codigo_barras"):
            codigos_planilha.setdefault(r["codigo_barras"], r["id"])
        chave = _chave_sem_codigo(r.get("beneficiario"), r.get("numero_documento"), r.get("valor"), r.get("vencimento"))
        if chave:
            chaves_planilha.setdefault(chave, r["id"])
    codigos_arquivo, chaves_arquivo = {}, {}
    itens = []

    for numero, valores in tabela["linhas"]:
        bruto = {campo: (valores[i] if (i := mapeamento.get(campo)) is not None and i < len(valores) else None)
                 for campo in CHAVES_CAMPOS}
        linha_txt = texto(bruto["linha_digitavel"])
        beneficiario = texto(bruto["beneficiario"])[:250]
        vencimento_planilha = ler_data(bruto["vencimento"])
        if not (linha_txt or beneficiario or vencimento_planilha):
            continue  # linha vazia, de total ou subtotal
        erros, avisos = [], []
        dados = {
            "linha_digitavel": "", "codigo_barras": "", "emissor": "", "tipo": None,
            "beneficiario": beneficiario,
            "documento_beneficiario": formatar_documento(texto(bruto["documento_beneficiario"]))[:30],
            "numero_documento": texto(bruto["numero_documento"])[:60],
            "observacoes": texto(bruto["observacoes"])[:1000],
            "centro_custo": (texto(bruto["centro_custo"]) or centro_padrao)[:60],
            "arquivo_nome": f"Importado de {nome_arquivo}"[:200] if nome_arquivo else "",
        }

        resultado = None
        if linha_txt:
            digitos = somente_digitos(linha_txt)
            resultado = analisar(digitos, referencia=vencimento_planilha or hoje) if digitos else None
            if resultado is None:
                erros.append(f"Linha digitável com {len(digitos)} dígitos (o esperado é 47, 48 ou 44)." if digitos
                             else "A coluna da linha digitável não tem números.")
            elif not resultado.valido:
                erros.append("Linha digitável inválida: " + "; ".join(resultado.erros) + ".")
            else:
                dados.update(linha_digitavel=resultado.linha_formatada, codigo_barras=resultado.codigo_barras,
                             emissor=resultado.emissor, tipo=resultado.tipo)
        else:
            avisos.append("Sem linha digitável: será lançado só com os dados da planilha.")
        codigo_ok = resultado is not None and resultado.valido

        valor_planilha = ler_valor(bruto["valor"])
        valor_codigo = float(resultado.valor) if codigo_ok and resultado.valor else None
        valor = valor_planilha or valor_codigo
        if valor_planilha and valor_codigo and abs(valor_planilha - valor_codigo) >= 0.01:
            avisos.append(f"Valor da planilha ({_brl(valor_planilha)}) diferente do código de barras ({_brl(valor_codigo)}).")
        if not valor:
            erros.append("Sem valor." if _vazio(bruto["valor"]) else f"Valor não reconhecido: “{texto(bruto['valor'])}”.")

        vencimento_codigo = resultado.vencimento if codigo_ok and not resultado.vencimento_estimado else None
        vencimento = vencimento_planilha or vencimento_codigo
        if vencimento_planilha and vencimento_codigo and vencimento_planilha != vencimento_codigo:
            avisos.append(f"Vencimento da planilha ({vencimento_planilha:%d/%m/%Y}) diferente do código de barras "
                          f"({vencimento_codigo:%d/%m/%Y}).")
        if not vencimento:
            erros.append("Sem vencimento." if _vazio(bruto["vencimento"])
                         else f"Vencimento não reconhecido: “{texto(bruto['vencimento'])}”.")
        if not beneficiario:
            avisos.append("Sem beneficiário.")

        data_pagamento = ler_data(bruto["data_pagamento"])
        status = ler_status(bruto["status"], "Pago" if data_pagamento else status_padrao)
        if status == "Pago" and not data_pagamento:
            data_pagamento = vencimento
            if vencimento:
                avisos.append("Pago sem data de pagamento: foi usada a data de vencimento.")
        if status != "Pago":
            data_pagamento = None
        dados.update(valor=valor, vencimento=vencimento, status=status, data_pagamento=data_pagamento)

        estado, mensagem = None, None
        if not erros:
            codigo = dados["codigo_barras"]
            chave = _chave_sem_codigo(beneficiario, dados["numero_documento"], valor, vencimento)
            existente = (codigos_planilha.get(codigo) if codigo else None) or (chaves_planilha.get(chave) if chave else None)
            anterior = (codigos_arquivo.get(codigo) if codigo else None) or (chaves_arquivo.get(chave) if chave else None)
            if existente is not None:
                estado = "duplicado"
                mensagem = f"Já lançado na planilha (ID {existente})." if existente > 0 else "Já lançado (aguardando gravação na planilha)."
            elif anterior is not None:
                estado, mensagem = "repetido", f"Repetido: igual à linha {anterior} do arquivo."
            if codigo:
                codigos_arquivo.setdefault(codigo, numero)
            if chave:
                chaves_arquivo.setdefault(chave, numero)
        estado = "erro" if erros else estado or ("aviso" if avisos else "pronto")
        itens.append({"indice": len(itens), "linha": numero, "estado": estado,
                      "mensagens": ([mensagem] if mensagem else []) + erros + avisos, "dados": dados})
    return itens


def serializar_item(item: dict) -> dict:
    dados = {k: (v.isoformat() if isinstance(v, date) else v) for k, v in item["dados"].items() if k != "arquivo_nome"}
    return {**item, "dados": dados}


def resumir(itens: list[dict]) -> dict:
    resumo = {estado: 0 for estado in ("pronto", "aviso", "erro", "duplicado", "repetido")}
    valor = 0.0
    for item in itens:
        resumo[item["estado"]] += 1
        if item["estado"] in ("pronto", "aviso"):
            valor += item["dados"]["valor"] or 0
    return {**resumo, "total": len(itens), "valor_importavel": round(valor, 2)}


# --------------------------------------------------------------------------- arquivos enviados e modelos


class ArmazemImportacoes:
    """Guarda por algumas horas o arquivo enviado, para a prévia e a importação sem reenviar."""

    def __init__(self, pasta: Path, horas: int = HORAS_GUARDAR_ARQUIVOS):
        self.pasta = Path(pasta)
        self.horas = horas
        self._lock = threading.Lock()
        self._cache: OrderedDict[str, tuple[str, OrderedDict]] = OrderedDict()

    def _caminhos(self, id_: str) -> tuple[Path, Path]:
        if not re.fullmatch(r"[0-9a-f]{32}", id_ or ""):
            raise KeyError(id_)
        return self.pasta / f"{id_}.bin", self.pasta / f"{id_}.json"

    def salvar(self, dados: bytes, nome: str) -> str:
        self.pasta.mkdir(parents=True, exist_ok=True)
        self.limpar()
        id_ = uuid.uuid4().hex
        arquivo, meta = self._caminhos(id_)
        arquivo.write_bytes(dados)
        meta.write_text(json.dumps({"nome": nome, "criado": time.time()}, ensure_ascii=False), encoding="utf-8")
        return id_

    def nome(self, id_: str) -> str:
        _, meta = self._caminhos(id_)
        try:
            return json.loads(meta.read_text(encoding="utf-8")).get("nome") or "arquivo"
        except (OSError, json.JSONDecodeError) as exc:
            raise KeyError(id_) from exc

    def planilhas(self, id_: str) -> tuple[str, OrderedDict]:
        arquivo, _ = self._caminhos(id_)
        with self._lock:
            if id_ in self._cache:
                self._cache.move_to_end(id_)
                return self._cache[id_]
        if not arquivo.exists():
            raise KeyError(id_)
        resultado = ler_planilhas(arquivo.read_bytes(), self.nome(id_))
        with self._lock:
            self._cache[id_] = resultado
            while len(self._cache) > 4:
                self._cache.popitem(last=False)
        return resultado

    def remover(self, id_: str) -> None:
        for caminho in self._caminhos(id_):
            caminho.unlink(missing_ok=True)
        with self._lock:
            self._cache.pop(id_, None)

    def limpar(self) -> int:
        if not self.pasta.exists():
            return 0
        limite = time.time() - self.horas * 3600
        removidos = 0
        for caminho in self.pasta.iterdir():
            try:
                if caminho.is_file() and caminho.stat().st_mtime < limite:
                    caminho.unlink()
                    removidos += 1
            except OSError:
                continue
        return removidos


class ModelosMapeamento:
    """Lembra quais colunas foram usadas em cada tipo de arquivo (identificado pelos títulos das colunas)."""

    def __init__(self, caminho: Path, limite: int = 50):
        self.caminho = Path(caminho)
        self.limite = limite
        self._lock = threading.Lock()

    @staticmethod
    def _assinatura(colunas: list[str]) -> str:
        return "|".join(normalizar_titulo(c) for c in colunas)

    def _ler(self) -> dict:
        try:
            return json.loads(self.caminho.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return {}

    def obter(self, colunas: list[str]) -> dict[str, int | None] | None:
        modelo = self._ler().get(self._assinatura(colunas))
        if not modelo:
            return None
        posicoes = {normalizar_titulo(t): i for i, t in enumerate(colunas)}
        return {campo: posicoes.get(normalizar_titulo(modelo["colunas"][campo])) if modelo["colunas"].get(campo) else None
                for campo in CHAVES_CAMPOS}

    def salvar(self, colunas: list[str], mapeamento: dict) -> None:
        with self._lock:
            modelos = self._ler()
            modelos[self._assinatura(colunas)] = {
                "colunas": {campo: colunas[i] for campo, i in mapeamento.items() if i is not None and i < len(colunas)},
                "usado_em": datetime.now().isoformat(timespec="seconds"),
            }
            if len(modelos) > self.limite:
                antigos = sorted(modelos, key=lambda k: modelos[k].get("usado_em", ""))
                for chave in antigos[: len(modelos) - self.limite]:
                    modelos.pop(chave)
            self.caminho.parent.mkdir(parents=True, exist_ok=True)
            temporario = self.caminho.with_suffix(".tmp")
            temporario.write_text(json.dumps(modelos, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(temporario, self.caminho)


def gerar_modelo() -> bytes:
    """Planilha modelo para quem pode escolher as colunas exportadas pelo outro sistema."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Boletos"
    colunas = [("Linha digitável", 56, "@"), ("Beneficiário", 34, None), ("CPF/CNPJ", 20, "@"), ("Nº do documento", 16, "@"),
               ("Vencimento", 13, "dd/mm/yyyy"), ("Valor", 14, '"R$" #,##0.00'), ("Situação", 12, None),
               ("Data de pagamento", 17, "dd/mm/yyyy"), ("Centro de custo", 24, None), ("Observações", 32, None)]
    exemplo = ["", "Fornecedor Exemplo Ltda", "02.359.273/0001-52", "NF 1001", date.today() + timedelta(days=10), 150.0,
               "Em aberto", None, "Administrativo", "Linha de exemplo: apague antes de importar"]
    for i, (titulo, largura, formato) in enumerate(colunas, start=1):
        celula = ws.cell(row=1, column=i, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F3D3E")
        celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largura
        valor = ws.cell(row=2, column=i, value=exemplo[i - 1])
        if formato:
            valor.number_format = formato
    ws.freeze_panes = "A2"
    ajuda = wb.create_sheet("Como preencher")
    ajuda.column_dimensions["A"].width = 110
    for linha in [
        "Como preencher a planilha de importação",
        "",
        "• Uma linha por boleto. Os títulos da primeira linha podem ter outros nomes: o sistema reconhece e você confere antes de importar.",
        "• Linha digitável: formate a coluna como Texto, senão o Excel arredonda os números.",
        "• Sem linha digitável, informe ao menos beneficiário, vencimento e valor.",
        "• Situação: Pago, Em aberto ou Cancelado. Com data de pagamento preenchida, o boleto entra como pago.",
        "• Centro de custo: onde o dinheiro foi gasto (obra, setor, veículo…). Nomes novos são cadastrados automaticamente.",
        "• Boletos já lançados na planilha são identificados e ficam de fora da importação.",
    ]:
        ajuda.append([linha])
    ajuda["A1"].font = Font(bold=True, size=13, color="0F3D3E")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
