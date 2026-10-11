"""API do caixa: lançamentos de entrada e saída, resumo por período, relatório em Excel e o bot do caixa."""

from __future__ import annotations

import io
from datetime import date
from typing import Literal

from fastapi import APIRouter, Request
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

from .. import contexto as ctx
from ..caixa import servico
from .comum import XLSX, anexo, autor, erro

router = APIRouter(prefix="/api/caixa", tags=["caixa"])
Periodo = Literal["dia", "semana", "mes", "ano"]


class LancamentoEntrada(BaseModel):
    movimento: Literal["Entrada", "Saída"] | None = None
    valor: float | None = Field(None, gt=0, le=1e10)
    data: date | None = None
    hora: str | None = Field(None, max_length=5)
    descricao: str | None = Field(None, max_length=200)
    local: str | None = Field(None, max_length=120)
    centro_custo: str | None = Field(None, max_length=60)
    pessoa: str | None = Field(None, max_length=60)
    observacoes: str | None = Field(None, max_length=300)


class ConfirmacaoCaixa(LancamentoEntrada):
    pass


class ConfigBotCaixa(BaseModel):
    token: str | None = Field(None, max_length=200)
    ativo: bool | None = None


def _falha(exc: Exception):
    return erro(422, "caixa", str(exc))


# --------------------------------------------------------------------------- lançamentos e resumo


@router.get("")
def resumo(periodo: Periodo = "mes", data: date | None = None):
    return servico.resumo(periodo, data)


@router.post("", status_code=201)
def lancar(entrada: LancamentoEntrada, request: Request):
    try:
        return servico.lancar(entrada.model_dump(exclude_none=True), autor(request))
    except servico.ErroCaixa as exc:
        raise _falha(exc) from exc


@router.put("/{id_}")
def editar(id_: int, entrada: LancamentoEntrada, request: Request):
    try:
        return servico.atualizar(id_, entrada.model_dump(exclude_unset=True), autor(request))
    except servico.ErroCaixa as exc:
        raise _falha(exc) from exc


@router.post("/{id_}/cancelar")
def cancelar(id_: int, request: Request):
    return servico.cancelar(id_, autor(request))


@router.get("/{id_}/historico")
def historico(id_: int):
    registro = ctx.repositorio_caixa.obter(id_)
    return {"eventos": ctx.planilha.historico_do_registro(id_, "caixa", registro.get("data_lancamento"))}


@router.get("/relatorio/arquivo")
def relatorio_excel(periodo: Periodo = "mes", data: date | None = None):
    """Excel do período: resumo, onde mais se gastou e todos os lançamentos."""
    r = servico.resumo(periodo, data)
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    cor = PatternFill("solid", fgColor="8A5200")
    ws["A1"] = f"Caixa · {r['periodo']['rotulo']}"
    ws["A1"].font = Font(bold=True, size=15, color="8A5200")
    linhas = [("Entradas", r["entradas"]["total"], r["entradas"]["quantidade"]),
              ("Saídas", r["saidas"]["total"], r["saidas"]["quantidade"]),
              ("Saldo do período", r["saldo_periodo"], None), ("Saldo do caixa no fim do período", r["saldo_caixa"], None)]
    ws.append([])
    for nome, valor, quantidade in linhas:
        ws.append([nome, valor, f"{quantidade} lançamento(s)" if quantidade is not None else None])
        ws.cell(row=ws.max_row, column=2).number_format = '"R$" #,##0.00'
    for titulo, chave in (("Saídas por centro de custo", "por_centro"), ("Onde mais se gastou", "por_local"),
                          ("Saídas por pessoa", "por_pessoa")):
        ws.append([])
        ws.append([titulo, "Total", "Quantidade", "%"])
        for celula in ws[ws.max_row]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = cor
        for item in r[chave]:
            ws.append([item["nome"], item["total"], item["quantidade"], item["percentual"] / 100])
            ws.cell(row=ws.max_row, column=2).number_format = '"R$" #,##0.00'
            ws.cell(row=ws.max_row, column=4).number_format = "0.0%"
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18

    lancamentos = wb.create_sheet("Lançamentos")
    colunas = [("Data", 12), ("Hora", 8), ("Entrada/Saída", 13), ("Valor", 14), ("O que foi pedido / comprado", 40),
               ("Onde foi gasto", 28), ("Centro de Custo", 24), ("Quem", 18), ("Lançado por", 18)]
    lancamentos.append([c for c, _ in colunas])
    for indice, (_, largura) in enumerate(colunas, start=1):
        celula = lancamentos.cell(row=1, column=indice)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = cor
        celula.alignment = Alignment(horizontal="center")
        lancamentos.column_dimensions[celula.column_letter].width = largura
    for item in sorted(r["lancamentos"], key=lambda i: (str(i.get("data")), i.get("hora") or "")):
        lancamentos.append([item.get("data"), item.get("hora") or "", item.get("movimento"), item.get("valor"),
                            item.get("descricao") or "", item.get("local") or "", item.get("centro_custo") or "",
                            item.get("pessoa") or "", item.get("lancado_por") or ""])
        lancamentos.cell(row=lancamentos.max_row, column=1).number_format = "dd/mm/yyyy"
        lancamentos.cell(row=lancamentos.max_row, column=4).number_format = '"R$" #,##0.00'
    lancamentos.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    inicio = r["periodo"]["inicio"]
    return anexo(buffer.getvalue(), XLSX, f"caixa-{periodo}-{inicio:%Y-%m-%d}.xlsx")


# --------------------------------------------------------------------------- bot do caixa (conferência)


@router.get("/bot/entrada")
def conferencia(incluir_concluidos: bool = False):
    return {"itens": ctx.caixa_bot.listar(incluir_concluidos), "pendentes": ctx.caixa_bot.contagem(),
            "config": ctx.caixa_bot.config_publica()}


@router.post("/bot/entrada/{id_}/lancar", status_code=201)
def lancar_da_conferencia(id_: str, entrada: ConfirmacaoCaixa, request: Request):
    try:
        return ctx.caixa_bot.confirmar(id_, entrada.model_dump(exclude_none=True), autor(request))
    except servico.ErroCaixa as exc:
        raise _falha(exc) from exc


@router.post("/bot/entrada/{id_}/descartar")
def descartar_da_conferencia(id_: str):
    if not ctx.caixa_bot.descartar(id_):
        raise erro(404, "nao_encontrado", "Este comprovante não está mais esperando revisão.")
    return {"ok": True}


@router.get("/bot/config")
def obter_config_bot():
    return ctx.caixa_bot.config_publica()


@router.post("/bot/config")
def salvar_config_bot(entrada: ConfigBotCaixa):
    try:
        return ctx.caixa_bot.configurar(entrada.model_dump(exclude_none=True))
    except ValueError as exc:
        raise erro(422, "config_invalida", str(exc)) from exc


@router.post("/bot/verificar")
def verificar_bot():
    try:
        return ctx.caixa_bot.verificar_agora()
    except ValueError as exc:
        raise erro(422, "bot_caixa", str(exc)) from exc
