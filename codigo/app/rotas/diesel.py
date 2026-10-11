"""API do diesel: máquinas, compras, abastecimentos, saldo, relatórios (tela, Excel e Telegram) e o bot do diesel."""

from __future__ import annotations

import io
from datetime import date
from typing import Literal

from fastapi import APIRouter, File, Request, UploadFile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

from .. import config
from .. import contexto as ctx
from ..diesel import leitura, precos, servico
from ..diesel import relatorio as relatorio_diesel
from ..diesel.modelos import COMBUSTIVEIS, MEDIDORES, TIPOS_MAQUINA
from .comum import XLSX, anexo, autor, erro

router = APIRouter(prefix="/api/diesel", tags=["diesel"])
Periodo = Literal["semana", "quinzena", "mes", "ano"]
Aba = Literal["compras", "abastecimentos"]
Combustivel = Literal[COMBUSTIVEIS]  # type: ignore[valid-type]


class MaquinaEntrada(BaseModel):
    nome: str | None = Field(None, max_length=60)
    tipo: Literal[TIPOS_MAQUINA] | None = None  # type: ignore[valid-type]
    placa: str | None = Field(None, max_length=16)
    medidor: Literal[MEDIDORES] | None = None  # type: ignore[valid-type]
    combustivel: Combustivel | Literal[""] | None = None
    status: Literal["Ativa", "Inativa"] | None = None
    observacoes: str | None = Field(None, max_length=300)


class CompraEntrada(BaseModel):
    data: date | None = None
    posto: str | None = Field(None, max_length=80)
    combustivel: Combustivel | None = None
    litros: float | None = Field(None, gt=0, le=1e6)
    valor: float | None = Field(None, ge=0, le=1e9)
    preco_litro: float | None = Field(None, ge=0, le=1000)
    forma: str | None = Field(None, max_length=20)
    documento: str | None = Field(None, max_length=30)
    observacoes: str | None = Field(None, max_length=300)
    arquivo_id: str | None = None


class AbastecimentoEntrada(BaseModel):
    data: date | None = None
    hora: str | None = Field(None, max_length=5)
    maquina: str | None = Field(None, max_length=60)
    combustivel: Combustivel | None = None
    litros: float | None = Field(None, gt=0, le=1e5)
    preco_litro: float | None = Field(None, ge=0, le=1000)
    valor: float | None = Field(None, ge=0, le=1e8)
    medidor: float | None = Field(None, ge=0, le=1e8)
    posto: str | None = Field(None, max_length=80)
    documento: str | None = Field(None, max_length=30)
    pessoa: str | None = Field(None, max_length=60)
    observacoes: str | None = Field(None, max_length=300)
    arquivo_id: str | None = None


class EdicaoEntrada(AbastecimentoEntrada):
    """Editar uma compra ou um abastecimento (cada aba aproveita só os seus campos)."""
    forma: str | None = Field(None, max_length=20)


class LinhaCanhoto(BaseModel):
    """Uma linha do canhoto: um combustível para uma máquina."""
    maquina: str | None = Field(None, max_length=60)
    combustivel: Combustivel | None = None
    litros: float | None = Field(None, ge=0, le=1e5)
    preco_litro: float | None = Field(None, ge=0, le=1000)
    valor: float | None = Field(None, ge=0, le=1e8)
    medidor: float | None = Field(None, ge=0, le=1e8)


class CanhotoEntrada(BaseModel):
    """Canhoto lançado pela tela: o que vale para todas as linhas + as linhas."""
    data: date | None = None
    hora: str | None = Field(None, max_length=5)
    posto: str | None = Field(None, max_length=80)
    documento: str | None = Field(None, max_length=30)
    pessoa: str | None = Field(None, max_length=60)
    observacoes: str | None = Field(None, max_length=300)
    arquivo_id: str | None = None
    itens: list[LinhaCanhoto] = Field(..., min_length=1, max_length=10)


class ConfirmacaoEntrada(EdicaoEntrada):
    tipo: Literal["compra", "canhoto"] | None = None
    itens: list[LinhaCanhoto] | None = Field(None, max_length=10)


class PrecosEntrada(BaseModel):
    precos: dict[Combustivel, float]


class ConfigBotDiesel(BaseModel):
    token: str | None = Field(None, max_length=200)
    ativo: bool | None = None
    gestores: list[str] | None = Field(None, max_length=50)
    alerta_valor: float | None = Field(None, ge=0, le=10000000)
    hora_relatorio: str | None = Field(None, pattern=r"^\d{2}:\d{2}$")
    relatorios: dict[Literal["semana", "quinzena", "mes"], bool] | None = None
    nivel_relatorio: Literal["simples", "detalhado", "completo"] | None = None


def _falha(exc: Exception):
    return erro(422, "diesel", str(exc))


def _arquivo(arquivo_id: str | None) -> dict | None:
    meta = ctx.arquivos.meta(arquivo_id) if arquivo_id else None
    return {"caminho": ctx.arquivos.caminho_relativo(arquivo_id) or "", "nome": meta["nome"]} if meta else None


# --------------------------------------------------------------------------- resumo, compras e abastecimentos


@router.get("")
def resumo(periodo: Periodo = "semana", data: date | None = None):
    r = servico.resumo(periodo, data)
    r["pendentes"] = ctx.diesel_bot.contagem() if ctx.diesel_bot else 0
    r["precos"] = precos.publico()
    r["alerta_valor"] = ctx.diesel_bot.alerta_valor() if ctx.diesel_bot else servico.ALERTA_PADRAO
    return r


@router.post("/arquivo", status_code=201)
async def guardar_arquivo(arquivo: UploadFile = File(...)):
    """Guarda o comprovante ou o canhoto escolhido na tela; devolve o id para o lançamento."""
    dados = await arquivo.read(config.MAX_UPLOAD_MB * 1024 * 1024 + 1)
    if not dados or len(dados) > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise erro(422, "arquivo_invalido", f"Arquivo vazio ou maior que {config.MAX_UPLOAD_MB} MB.")
    meta = ctx.arquivos.salvar(dados, arquivo.filename or "comprovante")
    return {"id": meta["id"], "nome": meta["nome"]}


@router.post("/analisar")
async def analisar_arquivo(arquivo: UploadFile = File(...)):
    """Lê o canhoto ou o comprovante enviado pela tela (sem passar pelo bot): combustíveis, litros, preços e totais."""
    dados = await arquivo.read(config.MAX_UPLOAD_MB * 1024 * 1024 + 1)
    if not dados or len(dados) > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise erro(422, "arquivo_invalido", f"Arquivo vazio ou maior que {config.MAX_UPLOAD_MB} MB.")
    meta = ctx.arquivos.salvar(dados, arquivo.filename or "canhoto")
    try:
        lido = leitura.analisar(dados, arquivo.filename or "", "")
    except Exception as exc:  # a leitura falhou: a pessoa preenche na mão, com o arquivo já guardado
        lido = {"tipo": "", "itens": [], "erro": f"Não consegui ler o arquivo: {exc}"}
    lido.pop("valores", None)
    maquina = servico.maquina_por_nome(lido.get("placa") or "") if lido.get("placa") else None
    return {**lido, "maquina_da_placa": (maquina or {}).get("nome") or "",
            "arquivo": {"id": meta["id"], "nome": meta["nome"]}}


@router.post("/canhoto", status_code=201)
def lancar_canhoto(entrada: CanhotoEntrada, request: Request):
    """Canhoto com uma ou mais linhas (combustíveis diferentes, máquinas diferentes): um abastecimento por linha."""
    comum = entrada.model_dump(exclude_none=True, exclude={"itens", "arquivo_id"})
    try:
        registros = servico.lancar_canhoto([i.model_dump(exclude_none=True) for i in entrada.itens], comum, autor(request),
                                           _arquivo(entrada.arquivo_id))
    except servico.ErroDiesel as exc:
        raise _falha(exc) from exc
    return {"abastecimentos": registros, "saldo": servico.saldo()}


@router.get("/precos")
def obter_precos():
    return {"precos": precos.publico(), "combustiveis": list(COMBUSTIVEIS)}


@router.post("/precos")
def salvar_precos(entrada: PrecosEntrada, request: Request):
    try:
        precos.salvar(entrada.precos, autor(request))
    except ValueError as exc:
        raise _falha(exc) from exc
    return obter_precos()


@router.post("/compras", status_code=201)
def lancar_compra(entrada: CompraEntrada, request: Request):
    try:
        return servico.lancar_compra(entrada.model_dump(exclude_none=True), autor(request), _arquivo(entrada.arquivo_id))
    except servico.ErroDiesel as exc:
        raise _falha(exc) from exc


@router.post("/abastecimentos", status_code=201)
def lancar_abastecimento(entrada: AbastecimentoEntrada, request: Request):
    try:
        return servico.lancar_abastecimento(entrada.model_dump(exclude_none=True), autor(request),
                                            _arquivo(entrada.arquivo_id))
    except servico.ErroDiesel as exc:
        raise _falha(exc) from exc


# --------------------------------------------------------------------------- máquinas


@router.get("/maquinas")
def listar_maquinas():
    return {"maquinas": [{**m, "media": servico.media_da_maquina(m["nome"])} for m in servico.maquinas()]}


@router.post("/maquinas", status_code=201)
def criar_maquina(entrada: MaquinaEntrada, request: Request):
    try:
        return servico.salvar_maquina(entrada.model_dump(exclude_none=True), autor(request))
    except servico.ErroDiesel as exc:
        raise _falha(exc) from exc


@router.put("/maquinas/{id_}")
def editar_maquina(id_: int, entrada: MaquinaEntrada, request: Request):
    try:
        return servico.salvar_maquina(entrada.model_dump(exclude_unset=True), autor(request), id_)
    except servico.ErroDiesel as exc:
        raise _falha(exc) from exc


# as rotas de máquinas vêm antes: "/{aba}/{id_}" também casaria com "/maquinas/3"


@router.put("/{aba}/{id_}")
def editar(aba: Aba, id_: int, entrada: EdicaoEntrada, request: Request):
    try:
        return servico.atualizar(aba, id_, entrada.model_dump(exclude_unset=True), autor(request))
    except servico.ErroDiesel as exc:
        raise _falha(exc) from exc


@router.post("/{aba}/{id_}/cancelar")
def cancelar(aba: Aba, id_: int, request: Request):
    return servico.cancelar(aba, id_, autor(request))


@router.get("/{aba}/{id_}/historico")
def historico(aba: Aba, id_: int):
    registro = ctx.diesel.obter(id_, aba)
    return {"eventos": ctx.diesel.historico_do_registro(id_, aba, registro.get("data_lancamento"))}


# --------------------------------------------------------------------------- relatórios


Nivel = Literal["simples", "detalhado", "completo"]


def _alerta_valor() -> float:
    return ctx.diesel_bot.alerta_valor() if ctx.diesel_bot else servico.ALERTA_PADRAO


@router.get("/relatorio")
def relatorio(periodo: Periodo = "semana", data: date | None = None):
    """O relatório do período com tudo o que os três níveis (simples, detalhado e completo) mostram."""
    r = relatorio_diesel.analise(periodo, data, _alerta_valor())
    r["pendentes"] = ctx.diesel_bot.contagem() if ctx.diesel_bot else 0
    r["precos"] = precos.publico()
    r["alerta_valor"] = _alerta_valor()
    return r


MOEDA = '"R$" #,##0.00'
LITROS = '#,##0.0 "L"'


@router.get("/relatorio/arquivo")
def relatorio_excel(periodo: Periodo = "semana", data: date | None = None):
    """Excel do período, sempre completo: resumo, máquinas, comparação, por dia, combustíveis, postos, pessoas,
    abastecimentos (com o consumo de cada trecho), compras e os pontos de atenção."""
    r = relatorio_diesel.analise(periodo, data, _alerta_valor())
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    cor = PatternFill("solid", fgColor="B45309")
    branco = Font(bold=True, color="FFFFFF")
    ws["A1"] = f"Diesel · relatório {relatorio_diesel.TITULOS[periodo]} · {r['periodo']['rotulo']}"
    ws["A1"].font = Font(bold=True, size=15, color="B45309")
    ws.append([])
    previsao = r.get("previsao") or {}
    linhas_resumo = [
        ("Abastecido (R$)", r["abastecido"]["custo"], MOEDA),
        ("Abastecido (litros)", r["abastecido"]["litros"], LITROS),
        ("Abastecimentos", r["abastecido"]["quantidade"], "0"),
        ("Litros por dia", r["abastecido"]["por_dia"], LITROS),
        (f"Abastecido em {r['anterior']['rotulo']} (R$)", r["anterior"]["custo"], MOEDA),
        ("Variação do valor", (r["variacao"]["custo"] or 0) / 100 if r["variacao"]["custo"] is not None else "—", "0.0%"),
        ("Comprado (R$)", r["comprado"]["valor"], MOEDA),
        ("Comprado (litros)", r["comprado"]["litros"], LITROS),
        ("Preço médio do litro nas compras", r["preco_medio"], '"R$" #,##0.000'),
        ("Saldo no posto no fim do período (R$)", r["saldo"]["valor"], MOEDA),
    ]
    if previsao:
        linhas_resumo += [("Gasto médio por dia (últimos 30 dias)", previsao["por_dia"], MOEDA),
                          ("O saldo dura (dias)", previsao["dias"] if previsao["dias"] is not None else "—", "0")]
    for nome, valor, formato in linhas_resumo:
        ws.append([nome, valor])
        ws.cell(row=ws.max_row, column=2).number_format = formato
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 18

    def aba(titulo: str, colunas: list[tuple[str, object, int, str | None]], linhas: list[dict]) -> None:
        """``colunas``: (título, chave ou função, largura, formato)."""
        folha = wb.create_sheet(titulo)
        folha.append([c[0] for c in colunas])
        for indice, (_, _, largura, _) in enumerate(colunas, start=1):
            celula = folha.cell(row=1, column=indice)
            celula.font, celula.fill, celula.alignment = branco, cor, Alignment(horizontal="center")
            folha.column_dimensions[celula.column_letter].width = largura
        for linha in linhas:
            folha.append([chave(linha) if callable(chave) else linha.get(chave) for _, chave, _, _ in colunas])
            for indice, (_, _, _, formato) in enumerate(colunas, start=1):
                if formato:
                    folha.cell(row=folha.max_row, column=indice).number_format = formato
        folha.freeze_panes = "A2"

    pct = lambda chave: lambda l: (l.get(chave) or 0) / 100  # noqa: E731
    aba("Máquinas", [("Máquina", "maquina", 30, None), ("Tipo", "tipo", 12, None), ("Placa / Prefixo", "placa", 14, None),
                     ("Combustíveis", lambda m: ", ".join(m["combustiveis"]), 22, None), ("Litros", "litros", 12, LITROS),
                     ("Valor", "custo", 14, MOEDA), ("R$ por litro", "custo_por_litro", 12, '"R$" #,##0.000'),
                     ("% do total", pct("percentual"), 11, "0.0%"), ("Abastecimentos", "abastecimentos", 14, "0"),
                     ("Litros por dia", "por_dia", 13, LITROS), ("Média por abastecimento", "por_abastecimento", 20, LITROS),
                     ("Consumo (L/h ou km/L)", lambda m: m["consumo"]["rotulo"], 20, None),
                     (f"Litros em {r['anterior']['rotulo']}", "litros_anterior", 18, LITROS),
                     ("Variação", lambda m: m["variacao"] / 100 if m["variacao"] is not None else "—", 10, "0.0%"),
                     ("Normal (L/dia, 90 dias antes)", "media_normal", 22, LITROS),
                     ("Acima do normal", lambda m: m["acima_do_normal"] / 100 if m["acima_do_normal"] else "", 14, "0%"),
                     ("Último", "ultimo", 12, "dd/mm/yyyy")], r["por_maquina"])
    aba("Por dia" if periodo != "ano" else "Por mês",
        [("Dia" if periodo != "ano" else "Mês", "rotulo", 12, None), ("Litros", "litros", 12, LITROS),
         ("Valor", "valor", 14, MOEDA), ("Abastecimentos", "abastecimentos", 14, "0"),
         *((c, (lambda c: lambda l: l["por_combustivel"].get(c, 0))(c), 16, MOEDA)
           for c in sorted({c for s in r["serie"] for c in s["por_combustivel"]}))], r["serie"])
    aba("Combustíveis", [("Combustível", "combustivel", 16, None), ("Comprado", "comprado", 12, LITROS),
                         ("Comprado (R$)", "valor_comprado", 14, MOEDA), ("Abastecido", "abastecido", 12, LITROS),
                         ("Abastecimentos", "abastecimentos", 14, "0"), ("Valor abastecido", "custo", 16, MOEDA),
                         ("Preço médio", "preco_medio", 12, '"R$" #,##0.000')], r["por_combustivel"])
    for titulo, lista in (("Por posto", r["por_posto"]), ("Por pessoa", r["por_pessoa"])):
        aba(titulo, [("Nome", "nome", 32, None), ("Litros", "litros", 12, LITROS), ("Valor", "valor", 14, MOEDA),
                     ("% do valor", pct("percentual"), 11, "0.0%"), ("Abastecimentos", "abastecimentos", 14, "0")], lista)
    trecho = lambda chave: lambda a: (a.get("trecho") or {}).get(chave)  # noqa: E731
    aba("Abastecimentos", [("Data", "data", 12, "dd/mm/yyyy"), ("Hora", "hora", 8, None), ("Máquina", "maquina", 30, None),
                           ("Combustível", "combustivel", 14, None), ("Litros", "litros", 12, LITROS),
                           ("Preço/L", "preco_litro", 12, '"R$" #,##0.000'), ("Valor", "custo", 14, MOEDA),
                           ("Horímetro / Km", "medidor", 14, "#,##0.0"), ("Leitura anterior", trecho("anterior"), 15, "#,##0.0"),
                           ("Rodado (h ou km)", trecho("rodado"), 15, "#,##0.0"), ("Consumo do trecho", trecho("rotulo"), 16, None),
                           ("Pedido", "documento", 12, None), ("Posto", "posto", 24, None),
                           ("Quem abasteceu", "pessoa", 20, None), ("Lançado por", "lancado_por", 20, None)],
        sorted(r["abastecimentos"], key=lambda l: (str(l.get("data")), l.get("hora") or "")))
    aba("Compras", [("Data", "data", 12, "dd/mm/yyyy"), ("Posto", "posto", 30, None), ("Combustível", "combustivel", 14, None),
                    ("Litros", "litros", 12, LITROS), ("Valor", "valor", 14, MOEDA), ("Preço/L", "preco_litro", 12, '"R$" #,##0.000'),
                    ("Forma", "forma", 14, None), ("Nota / Cupom", "documento", 16, None),
                    ("Lançada por", "lancado_por", 20, None)], sorted(r["compras"], key=lambda l: str(l.get("data"))))
    aba("Pontos de atenção", [("Tipo", lambda a: {"erro": "Grave", "aviso": "Atenção", "info": "Informação"}[a["nivel"]], 12, None),
                              ("O que foi visto", "texto", 110, None)], r["alertas"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return anexo(buffer.getvalue(), XLSX, f"diesel-{periodo}-{r['periodo']['inicio']:%Y-%m-%d}.xlsx")


@router.get("/relatorio/texto")
def relatorio_texto(periodo: Periodo = "semana", data: date | None = None, nivel: Nivel = "simples"):
    return {"texto": relatorio_diesel.texto(periodo, data, nivel, _alerta_valor())}


@router.post("/relatorio/enviar")
def enviar_relatorio(periodo: Periodo = "semana", data: date | None = None, nivel: Nivel = "simples"):
    """Manda agora o relatório (no nível escolhido) pelo bot do diesel para os responsáveis."""
    bot = ctx.diesel_bot
    destinos = bot.gestores() if bot else []
    if not bot or not bot.token() or not destinos:
        raise erro(422, "sem_destino", "Configure o bot do diesel e escolha quem recebe os relatórios.")
    texto = relatorio_diesel.texto(periodo, data, nivel, _alerta_valor())
    for chat_id in destinos:
        bot._enviar(chat_id, texto)
    return {"enviado": len(destinos)}


# --------------------------------------------------------------------------- bot do diesel (conferência)


@router.get("/bot/entrada")
def conferencia(incluir_concluidos: bool = False):
    return {"itens": ctx.diesel_bot.listar(incluir_concluidos), "pendentes": ctx.diesel_bot.contagem(),
            "config": ctx.diesel_bot.config_publica()}


@router.post("/bot/entrada/{id_}/lancar", status_code=201)
def lancar_da_conferencia(id_: str, entrada: ConfirmacaoEntrada, request: Request):
    try:
        resultado = ctx.diesel_bot.confirmar(id_, entrada.model_dump(exclude_none=True), autor(request))
        resultado.pop("registros", None)
        return resultado
    except servico.ErroDiesel as exc:
        raise _falha(exc) from exc


@router.post("/bot/entrada/{id_}/descartar")
def descartar_da_conferencia(id_: str):
    if not ctx.diesel_bot.descartar(id_):
        raise erro(404, "nao_encontrado", "Este arquivo não está mais esperando conferência.")
    return {"ok": True}


@router.get("/bot/config")
def obter_config_bot():
    return ctx.diesel_bot.config_publica()


@router.post("/bot/config")
def salvar_config_bot(entrada: ConfigBotDiesel):
    try:
        return ctx.diesel_bot.configurar(entrada.model_dump(exclude_none=True))
    except ValueError as exc:
        raise erro(422, "config_invalida", str(exc)) from exc


@router.post("/bot/verificar")
def verificar_bot():
    try:
        return ctx.diesel_bot.verificar_agora()
    except ValueError as exc:
        raise erro(422, "bot_diesel", str(exc)) from exc
