"""Relatório de fechamento de uma empresa, no mesmo formato da planilha usada hoje.

Cabeçalho com a minha empresa e o CNPJ, "Fechamento <Empresa>" e as colunas DATA, MATERIAL, UNIDADE, QUANT.,
VALOR UNI., VALOR TOTAL, NOTA FISCAL e Nº PEDIDO, com o total no rodapé. Abaixo da tabela, o que já foi
recebido e quanto ainda falta. Sai em Excel e em PDF; o texto curto vai pelo Telegram.
"""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUNAS = [("DATA", 13), ("MATERIAL", 26), ("UNIDADE", 11), ("QUANT.", 10), ("VALOR UNI.", 13),
           ("VALOR TOTAL", 15), ("NOTA FISCAL", 13), ("Nº PEDIDO", 12)]
VERDE = "E2EFDA"
VERDE_FORTE = "C6E0B4"
AZUL = "1F4E79"
MOEDA = '"R$" #,##0.00'
LINHAS_EM_BRANCO = 7  # como na planilha usada hoje, algumas linhas ficam prontas para escrever à mão


def _moeda(valor) -> str:
    return "R$ " + f"{float(valor or 0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _data(valor) -> str:
    if isinstance(valor, datetime):
        valor = valor.date()
    return valor.strftime("%d/%m/%Y") if isinstance(valor, date) else ""


def cabecalho_minha_empresa(minha_empresa: dict) -> str:
    nome = (minha_empresa or {}).get("nome") or "Minha empresa"
    cnpj = (minha_empresa or {}).get("cnpj") or ""
    return f"{nome} {cnpj}".strip()


def nome_do_arquivo(f: dict, extensao: str) -> str:
    sem_acento = unicodedata.normalize("NFKD", f["empresa"]).encode("ascii", "ignore").decode("ascii")
    empresa = re.sub(r"[^A-Za-z0-9]+", "-", sem_acento).strip("-").lower() or "empresa"
    return f"fechamento-{empresa}-{f['competencia']}.{extensao}"


# --------------------------------------------------------------------------- Excel


def gerar_xlsx(f: dict, minha_empresa: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Fechamento {f['empresa']}"[:31]
    borda = Border(*(Side(style="thin", color="7F9DB9"),) * 4)
    ultima = get_column_letter(len(COLUNAS))

    ws.merge_cells(f"A1:{ultima}1")
    titulo = ws["A1"]
    titulo.value = cabecalho_minha_empresa(minha_empresa or {})
    titulo.font = Font(bold=True, size=11, color=AZUL)
    titulo.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(f"A2:{ultima}2")
    subtitulo = ws["A2"]
    subtitulo.value = f"Fechamento {f['empresa']}"
    subtitulo.font = Font(bold=True, size=11, color=AZUL)
    subtitulo.alignment = Alignment(horizontal="center", vertical="center")
    ws["A3"] = (f"Período de {_data(f['inicio'])} a {_data(f['fim'])} · gerado em {f['gerado_em']:%d/%m/%Y %H:%M}")
    ws["A3"].font = Font(size=9, color="666666")
    ws.merge_cells(f"A3:{ultima}3")
    ws["A3"].alignment = Alignment(horizontal="center")

    inicio = 5
    for coluna, (nome, largura) in enumerate(COLUNAS, start=1):
        celula = ws.cell(row=inicio, column=coluna, value=nome)
        celula.font = Font(bold=True)
        celula.fill = PatternFill("solid", fgColor=VERDE_FORTE)
        celula.alignment = Alignment(horizontal="center", vertical="center")
        celula.border = borda
        ws.column_dimensions[get_column_letter(coluna)].width = largura

    linha = inicio
    for item in f["linhas"]:
        linha += 1
        valores = [item.get("data"), item.get("material") or "", item.get("unidade") or "", item.get("quantidade"),
                   item.get("valor_unitario"), item.get("valor_total"), item.get("numero_nf") or "",
                   item.get("pedido") or ""]
        for coluna, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha, column=coluna, value=valor if valor != "" else None)
            celula.fill = PatternFill("solid", fgColor=VERDE)
            celula.border = borda
            if coluna == 1:
                celula.number_format = "dd/mm/yyyy"
                celula.alignment = Alignment(horizontal="center")
            elif coluna in (5, 6):
                celula.number_format = MOEDA
            elif coluna in (3, 4, 7, 8):
                celula.alignment = Alignment(horizontal="center")
    primeira_linha, ultima_linha = inicio + 1, linha
    for _ in range(LINHAS_EM_BRANCO):
        linha += 1
        for coluna in range(1, len(COLUNAS) + 1):
            celula = ws.cell(row=linha, column=coluna)
            celula.fill = PatternFill("solid", fgColor=VERDE)
            celula.border = borda
            if coluna in (5, 6):
                celula.number_format = MOEDA

    total = linha + 1
    rotulo = ws.cell(row=total, column=5, value="TOTAL")
    rotulo.font = Font(bold=True)
    rotulo.alignment = Alignment(horizontal="right")
    celula = ws.cell(row=total, column=6,
                     value=f"=SUM(F{primeira_linha}:F{linha})" if ultima_linha >= primeira_linha else 0)
    celula.font = Font(bold=True)
    celula.number_format = MOEDA
    celula.border = borda
    celula.fill = PatternFill("solid", fgColor=VERDE_FORTE)

    t = f["totais"]
    resumo = [("Total do fechamento", t["total"]), ("Já recebido", t["recebido"]), ("Falta receber", t["falta"])]
    if t.get("em_cheque"):
        resumo.append(("Em cheque (a compensar)", t["em_cheque"]))
    for i, (nome, valor) in enumerate(resumo, start=total + 2):
        ws.cell(row=i, column=5, value=nome).font = Font(bold=nome == "Falta receber")
        celula = ws.cell(row=i, column=6, value=valor)
        celula.number_format = MOEDA
        celula.font = Font(bold=nome == "Falta receber", color="B42318" if nome == "Falta receber" and valor else "000000")

    if f["recebimentos"]:
        linha_receb = total + len(resumo) + 4
        ws.cell(row=linha_receb, column=1, value="Recebimentos do período").font = Font(bold=True, color=AZUL)
        for coluna, (nome, largura) in enumerate([("Data", 13), ("Forma", 12), ("Pagador", 26), ("Notas", 16),
                                                  ("Valor", 15)], start=1):
            celula = ws.cell(row=linha_receb + 1, column=coluna, value=nome)
            celula.font = Font(bold=True)
            celula.fill = PatternFill("solid", fgColor=VERDE_FORTE)
            celula.border = borda
        for i, r in enumerate(f["recebimentos"], start=linha_receb + 2):
            valores = [r.get("data"), r.get("forma") or "", r.get("pagador") or "", r.get("notas") or "", r.get("valor")]
            for coluna, valor in enumerate(valores, start=1):
                celula = ws.cell(row=i, column=coluna, value=valor if valor != "" else None)
                celula.border = borda
                if coluna == 1:
                    celula.number_format = "dd/mm/yyyy"
                elif coluna == 5:
                    celula.number_format = MOEDA

    ws.sheet_view.showGridLines = False
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- PDF


def gerar_pdf(f: dict, minha_empresa: dict | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    estilos = getSampleStyleSheet()
    azul = colors.HexColor("#1F4E79")
    titulo = ParagraphStyle("titulo", parent=estilos["Title"], fontSize=13, textColor=azul, spaceAfter=0)
    sub = ParagraphStyle("sub", parent=estilos["Title"], fontSize=12, textColor=azul, spaceAfter=2)
    nota = ParagraphStyle("nota", parent=estilos["Normal"], fontSize=8, alignment=1, textColor=colors.HexColor("#666666"))
    celula = ParagraphStyle("celula", parent=estilos["Normal"], fontSize=8, leading=10)

    linhas = [[_data(i.get("data")), Paragraph(str(i.get("material") or ""), celula), i.get("unidade") or "",
               f"{float(i.get('quantidade') or 0):g}", _moeda(i.get("valor_unitario")), _moeda(i.get("valor_total")),
               i.get("numero_nf") or "", i.get("pedido") or ""] for i in f["linhas"]]
    if not linhas:
        linhas = [["", Paragraph("Nenhuma nota fiscal neste fechamento", celula), "", "", "", "", "", ""]]
    t = f["totais"]
    dados = [[c for c, _ in COLUNAS]] + linhas + [["", "", "", "", "TOTAL", _moeda(t["total"]), "", ""]]
    tabela = Table(dados, colWidths=[22, 46, 18, 16, 24, 27, 22, 22], repeatRows=1, hAlign="CENTER")
    tabela._argW = [w * mm for w in (22, 46, 18, 16, 24, 27, 22, 22)]
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#C6E0B4")),
        ("BACKGROUND", (0, 1), (-1, -2), colors.HexColor("#E2EFDA")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#C6E0B4")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#7F9DB9")),
        ("ALIGN", (2, 0), (-1, -1), "CENTER"),
        ("ALIGN", (4, 1), (5, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    resumo = [["Total do fechamento", _moeda(t["total"])], ["Já recebido", _moeda(t["recebido"])],
              ["Falta receber", _moeda(t["falta"])]]
    if t.get("em_cheque"):
        resumo.append(["Em cheque (a compensar)", _moeda(t["em_cheque"])])
    tabela_resumo = Table(resumo, colWidths=[60 * mm, 35 * mm], hAlign="CENTER")
    tabela_resumo.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("FONTNAME", (0, 2), (-1, 2), "Helvetica-Bold"),
        ("TEXTCOLOR", (1, 2), (1, 2), colors.HexColor("#B42318") if t["falta"] else colors.black),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D5DDE0")),
        ("ALIGN", (1, 0), (1, -1), "RIGHT"),
    ]))

    elementos = [
        Paragraph(cabecalho_minha_empresa(minha_empresa or {}), titulo),
        Paragraph(f"Fechamento {f['empresa']}", sub),
        Paragraph(f"Período de {_data(f['inicio'])} a {_data(f['fim'])} · gerado em {f['gerado_em']:%d/%m/%Y %H:%M}", nota),
        Spacer(1, 10), tabela, Spacer(1, 12), tabela_resumo,
    ]
    buffer = io.BytesIO()
    documento = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm,
                                  bottomMargin=14 * mm, title=f"Fechamento {f['empresa']} {f['competencia']}",
                                  author=cabecalho_minha_empresa(minha_empresa or {}))
    documento.build(elementos)
    return buffer.getvalue()


# --------------------------------------------------------------------------- texto (Telegram)


def resumo_telegram(f: dict, minha_empresa: dict | None = None) -> str:
    from html import escape

    t = f["totais"]
    linhas = [f"📄 <b>Fechamento {escape(f['empresa'])}</b> — {escape(f['titulo'])}",
              f"Período: {_data(f['inicio'])} a {_data(f['fim'])}",
              f"Notas: {t['quantidade']} lançamento(s)",
              f"Total: <b>{_moeda(t['total'])}</b>",
              f"Recebido: {_moeda(t['recebido'])}"]
    if t.get("em_cheque"):
        linhas.append(f"Em cheque (a compensar): {_moeda(t['em_cheque'])}")
    linhas.append(f"{'✅ Tudo recebido' if t['falta'] <= 0.01 else '🔴 Falta receber: <b>' + _moeda(t['falta']) + '</b>'}")
    return "\n".join(linhas)
