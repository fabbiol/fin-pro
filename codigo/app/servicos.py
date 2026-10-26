"""Regras de consulta sobre os registros da planilha: filtros, ordenação, painel, análise de gastos e exportação."""

from __future__ import annotations

import csv
import io
from collections import defaultdict
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .boleto.codigo import somente_digitos
from .boleto.texto import normalizar
from .pix import TIPOS_CHAVE, formatar_chave
from .planilha import COLUNAS, ROTULOS

MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
SITUACOES = ("vencido", "vence_hoje", "a_vencer", "pendente", "pago", "cancelado")
SEM_CENTRO = "Sem centro de custo"
OUTROS_CENTROS = "Outros centros"
CAMPOS_BUSCA_BOLETOS = ("beneficiario", "numero_documento", "documento_beneficiario", "observacoes", "emissor", "linha_digitavel", "centro_custo")
CAMPOS_DIGITOS_BOLETOS = ("codigo_barras", "linha_digitavel")
CAMPOS_BUSCA_PIX = ("recebedor", "chave", "descricao", "observacoes", "centro_custo", "txid")
CAMPOS_DIGITOS_PIX = ("chave",)


def situacao(registro: dict, hoje: date) -> str:
    status = registro.get("status")
    if status == "Pago":
        return "pago"
    if status == "Cancelado":
        return "cancelado"
    venc = registro.get("vencimento")
    if not isinstance(venc, date):
        return "pendente"
    if venc < hoje:
        return "vencido"
    if venc == hoje:
        return "vence_hoje"
    return "a_vencer"


def serializar(registro: dict, hoje: date | None = None) -> dict:
    hoje = hoje or date.today()
    saida = {}
    for chave, valor in registro.items():
        saida[chave] = valor.isoformat() if isinstance(valor, (date, datetime)) else valor
    venc = registro.get("vencimento")
    saida["situacao"] = situacao(registro, hoje)
    saida["dias_para_vencer"] = (venc - hoje).days if isinstance(venc, date) else None
    saida["linha_digitos"] = somente_digitos(registro.get("linha_digitavel"))
    return saida


def serializar_pix(registro: dict, hoje: date | None = None) -> dict:
    saida = serializar(registro, hoje)
    saida.pop("linha_digitos", None)
    saida["tipo_chave_rotulo"] = TIPOS_CHAVE.get(registro.get("tipo_chave"), registro.get("tipo_chave") or "")
    saida["chave_formatada"] = formatar_chave(registro.get("tipo_chave"), registro.get("chave"))
    return saida


def _valor(registro: dict) -> float:
    return float(registro.get("valor") or 0)


def _pago(registro: dict) -> float:
    """Quanto saiu de fato: o valor pago (com juros ou desconto) ou, sem ele, o valor do lançamento."""
    try:
        return float(registro.get("valor_pago") or 0) or _valor(registro)
    except (TypeError, ValueError):
        return _valor(registro)


def _juros(registro: dict) -> float:
    try:
        return round(float(registro.get("juros_multa") or 0), 2)
    except (TypeError, ValueError):
        return 0.0


def filtrar(
    registros: list[dict],
    *,
    q: str = "",
    situacoes: set[str] | None = None,
    campo_data: str = "vencimento",
    de: date | None = None,
    ate: date | None = None,
    hoje: date | None = None,
    centro: str = "",
    campos_busca: tuple[str, ...] = CAMPOS_BUSCA_BOLETOS,
    campos_digitos: tuple[str, ...] = CAMPOS_DIGITOS_BOLETOS,
) -> list[dict]:
    hoje = hoje or date.today()
    palavras = normalizar(q.strip()).split()
    termo_digitos = somente_digitos(q)
    centro_normalizado = (centro or "").strip().lower()
    resultado = []
    for r in registros:
        if situacoes:
            sit = situacao(r, hoje)
            aceitos = set(situacoes)
            if "pendente" in aceitos:
                aceitos |= {"vencido", "vence_hoje", "a_vencer"}
            if sit not in aceitos:
                continue
        if centro_normalizado:
            atual = (r.get("centro_custo") or "").strip().lower()
            if centro_normalizado == "__sem__":
                if atual:
                    continue
            elif atual != centro_normalizado:
                continue
        if de or ate:
            valor_data = r.get(campo_data)
            if isinstance(valor_data, datetime):
                valor_data = valor_data.date()
            if not isinstance(valor_data, date):
                continue
            if de and valor_data < de:
                continue
            if ate and valor_data > ate:
                continue
        if palavras:
            texto = normalizar(" ".join(str(r.get(k) or "") for k in campos_busca))
            achou = all(p in texto for p in palavras)
            if not achou and len(termo_digitos) >= 4:
                achou = any(termo_digitos in somente_digitos(str(r.get(k) or "")) for k in campos_digitos)
            if not achou:
                continue
        resultado.append(r)
    return resultado


def ordenar(registros: list[dict], campo: str = "vencimento", direcao: str = "asc") -> list[dict]:
    reverso = direcao == "desc"

    def chave(r):
        v = r.get(campo)
        if v is None or v == "":
            return (1, "")  # vazios sempre no fim
        if isinstance(v, datetime):
            return (0, v.timestamp())
        if isinstance(v, date):
            return (0, v.toordinal())
        if isinstance(v, (int, float)):
            return (0, v)
        return (0, normalizar(str(v)))

    preenchidos = [r for r in registros if chave(r)[0] == 0]
    vazios = [r for r in registros if chave(r)[0] == 1]
    return sorted(preenchidos, key=chave, reverse=reverso) + vazios


def resumir(registros: list[dict], hoje: date | None = None) -> dict:
    hoje = hoje or date.today()
    resumo = {"quantidade": len(registros), "valor_total": 0.0, "valor_pendente": 0.0, "valor_pago": 0.0, "valor_vencido": 0.0,
              "juros_multa": 0.0, "pagos_com_juros": 0}
    for r in registros:
        sit = situacao(r, hoje)
        if sit == "cancelado":
            continue
        v = _pago(r) if sit == "pago" else _valor(r)
        resumo["valor_total"] += v
        if sit == "pago":
            resumo["valor_pago"] += v
            juros = _juros(r)
            if juros:
                resumo["juros_multa"] += juros
                resumo["pagos_com_juros"] += 1
        else:
            resumo["valor_pendente"] += v
            if sit == "vencido":
                resumo["valor_vencido"] += v
    return {k: round(v, 2) if isinstance(v, float) else v for k, v in resumo.items()}


def _inicio_mes(d: date, deslocamento: int = 0) -> date:
    indice = d.month - 1 + deslocamento
    return date(d.year + indice // 12, indice % 12 + 1, 1)


def painel(registros: list[dict], hoje: date | None = None, pix: list[dict] | None = None) -> dict:
    """Indicadores, gráfico e listas do Painel, somando boletos e Pix (os Pix vêm marcados com ``tipo_pagamento``)."""
    hoje = hoje or date.today()
    registros = [*({**r, "tipo_pagamento": "boleto"} for r in registros),
                 *({**p, "tipo_pagamento": "pix", "beneficiario": p.get("recebedor") or p.get("beneficiario") or ""} for p in pix or [])]
    inicio_mes = _inicio_mes(hoje)
    fim_semana = hoje + timedelta(days=6)  # 7 dias contando hoje, como na Projeção

    k = defaultdict(float)
    for r in registros:
        sit = situacao(r, hoje)
        v = _valor(r)
        venc = r.get("vencimento")
        lancado = r.get("data_lancamento")
        pago_em = r.get("data_pagamento")
        if isinstance(lancado, datetime) and lancado.date() >= inicio_mes:
            k["lancados_mes_qtd"] += 1
            k["lancados_mes_valor"] += v
        if sit == "cancelado":
            continue
        k["total_qtd"] += 1
        if sit == "pago":
            if isinstance(pago_em, date) and pago_em >= inicio_mes:
                k["pagos_mes_qtd"] += 1
                k["pagos_mes_valor"] += _pago(r)
                if _juros(r):
                    k["juros_mes_qtd"] += 1
                    k["juros_mes_valor"] += _juros(r)
            continue
        k["pendentes_qtd"] += 1
        k["pendentes_valor"] += v
        if sit == "vencido":
            k["vencidos_qtd"] += 1
            k["vencidos_valor"] += v
            k["vencidos_pix_qtd" if r["tipo_pagamento"] == "pix" else "vencidos_boletos_qtd"] += 1
        elif isinstance(venc, date) and venc <= fim_semana:
            k["proximos7_qtd"] += 1
            k["proximos7_valor"] += v

    meses = [_inicio_mes(hoje, d) for d in range(-5, 7)]
    serie = {m: {"pago": 0.0, "vencido": 0.0, "a_vencer": 0.0} for m in meses}
    for r in registros:
        venc = r.get("vencimento")
        if not isinstance(venc, date):
            continue
        mes = date(venc.year, venc.month, 1)
        if mes not in serie:
            continue
        sit = situacao(r, hoje)
        if sit == "cancelado":
            continue
        chave = "pago" if sit == "pago" else "vencido" if sit == "vencido" else "a_vencer"
        serie[mes][chave] += _pago(r) if chave == "pago" else _valor(r)

    beneficiarios: dict[str, dict] = {}
    for r in registros:
        if situacao(r, hoje) == "cancelado" or not r.get("beneficiario"):
            continue
        chave = normalizar(r["beneficiario"]).strip()
        item = beneficiarios.setdefault(chave, {"beneficiario": r["beneficiario"], "quantidade": 0, "valor": 0.0})
        item["quantidade"] += 1
        item["valor"] += _valor(r)

    pendentes = [r for r in registros if situacao(r, hoje) in ("vencido", "vence_hoje", "a_vencer", "pendente")]
    com_data = [r for r in pendentes if isinstance(r.get("vencimento"), date)]
    recentes = sorted(
        (r for r in registros if isinstance(r.get("data_lancamento"), datetime)),
        key=lambda r: (r["data_lancamento"], r["id"]), reverse=True,
    )

    def item(r: dict) -> dict:
        saida = serializar_pix(r, hoje) if r["tipo_pagamento"] == "pix" else serializar(r, hoje)
        return {**saida, "tipo_pagamento": r["tipo_pagamento"], "beneficiario": r.get("beneficiario") or ""}

    return {
        "hoje": hoje.isoformat(),
        "indicadores": {chave: round(valor, 2) if "valor" in chave else int(valor) for chave, valor in {
            "total_qtd": k["total_qtd"],
            "pendentes_qtd": k["pendentes_qtd"], "pendentes_valor": k["pendentes_valor"],
            "vencidos_qtd": k["vencidos_qtd"], "vencidos_valor": k["vencidos_valor"],
            "vencidos_boletos_qtd": k["vencidos_boletos_qtd"], "vencidos_pix_qtd": k["vencidos_pix_qtd"],
            "proximos7_qtd": k["proximos7_qtd"], "proximos7_valor": k["proximos7_valor"],
            "pagos_mes_qtd": k["pagos_mes_qtd"], "pagos_mes_valor": k["pagos_mes_valor"],
            "juros_mes_qtd": k["juros_mes_qtd"], "juros_mes_valor": k["juros_mes_valor"],
            "lancados_mes_qtd": k["lancados_mes_qtd"], "lancados_mes_valor": k["lancados_mes_valor"],
        }.items()},
        "serie_mensal": [
            {"mes": m.isoformat(), "rotulo": f"{MESES[m.month - 1]}/{str(m.year)[2:]}", "atual": m == inicio_mes,
             **{chave: round(v, 2) for chave, v in valores.items()}}
            for m, valores in serie.items()
        ],
        # Todos os vencidos (boletos e Pix), do mais antigo para o mais recente
        "vencidos": [item(r) for r in sorted(
            (r for r in com_data if r["vencimento"] < hoje), key=lambda r: (r["vencimento"], r["id"]))],
        "proximos": [item(r) for r in sorted(
            (r for r in com_data if r["vencimento"] >= hoje), key=lambda r: r["vencimento"])[:8]],
        "top_beneficiarios": sorted(
            ({**b, "valor": round(b["valor"], 2)} for b in beneficiarios.values()), key=lambda b: b["valor"], reverse=True)[:6],
        "recentes": [item(r) for r in recentes[:6]],
    }


# --------------------------------------------------------------------------- análise de gastos


def analise_gastos(
    boletos: list[dict],
    pix: list[dict],
    caixa: list[dict] | None = None,
    *,
    de: date | None = None,
    ate: date | None = None,
    campo_data: str = "vencimento",
    situacao_filtro: str = "todos",
    tipo: str = "todos",
    hoje: date | None = None,
) -> dict:
    """Quanto foi gasto em cada centro de custo (boletos + Pix + dinheiro do caixa), por mês e por beneficiário."""
    hoje = hoje or date.today()
    itens = []
    if tipo in ("todos", "boleto"):
        itens += [("boleto", r) for r in boletos]
    if tipo in ("todos", "pix"):
        itens += [("pix", r) for r in pix]
    if tipo in ("todos", "dinheiro"):
        itens += [("dinheiro", r) for r in caixa or []]

    def gasto(r: dict) -> float:  # pago: o que saiu de fato; pendente: o valor do lançamento
        return _pago(r) if situacao(r, hoje) == "pago" else _valor(r)

    selecionados = []
    for origem, r in itens:
        sit = situacao(r, hoje)
        if sit == "cancelado":
            continue
        if situacao_filtro == "pagos" and sit != "pago":
            continue
        if situacao_filtro == "pendentes" and sit == "pago":
            continue
        data_ref = r.get(campo_data)
        if isinstance(data_ref, datetime):
            data_ref = data_ref.date()
        if de or ate:
            if not isinstance(data_ref, date):
                continue
            if (de and data_ref < de) or (ate and data_ref > ate):
                continue
        selecionados.append((origem, r, data_ref))

    def nome_centro(r):
        return (r.get("centro_custo") or "").strip() or SEM_CENTRO

    por_centro: dict[str, dict] = {}
    for origem, r, _ in selecionados:
        nome = nome_centro(r)
        item = por_centro.setdefault(nome.lower(), {"centro": nome, "total": 0.0, "boletos": 0.0, "pix": 0.0,
                                                    "dinheiro": 0.0, "quantidade": 0, "pagos": 0.0, "pendentes": 0.0})
        v = gasto(r)
        item["total"] += v
        item[{"boleto": "boletos", "pix": "pix"}.get(origem, "dinheiro")] += v
        item["quantidade"] += 1
        item["pagos" if situacao(r, hoje) == "pago" else "pendentes"] += v

    total = sum(i["total"] for i in por_centro.values())
    centros = sorted(por_centro.values(), key=lambda i: i["total"], reverse=True)
    for item in centros:
        item["percentual"] = round(item["total"] / total * 100, 1) if total else 0.0
        for chave in ("total", "boletos", "pix", "dinheiro", "pagos", "pendentes"):
            item[chave] = round(item[chave], 2)

    principais = [c["centro"] for c in centros[:6]]
    meses: dict[date, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for _, r, data_ref in selecionados:
        if isinstance(data_ref, date):
            meses[date(data_ref.year, data_ref.month, 1)][nome_centro(r)] += gasto(r)
    serie_mensal = []
    for mes in sorted(meses):
        valores = meses[mes]
        linha = {"mes": mes.isoformat(), "rotulo": f"{MESES[mes.month - 1]}/{str(mes.year)[2:]}",
                 "total": round(sum(valores.values()), 2), "centros": {}}
        outros = 0.0
        for nome, v in valores.items():
            if nome in principais:
                linha["centros"][nome] = round(v, 2)
            else:
                outros += v
        if outros:
            linha["centros"][OUTROS_CENTROS] = round(outros, 2)
        serie_mensal.append(linha)

    por_nome: dict[str, dict] = {}
    for origem, r, _ in selecionados:
        nome = (r.get("recebedor") if origem == "pix" else r.get("beneficiario")) or "Não informado"
        item = por_nome.setdefault(normalizar(nome).strip(), {"nome": nome, "total": 0.0, "quantidade": 0})
        item["total"] += gasto(r)
        item["quantidade"] += 1
    maiores = sorted(({**i, "total": round(i["total"], 2)} for i in por_nome.values()), key=lambda i: i["total"], reverse=True)[:10]

    return {
        "total": round(total, 2),
        "quantidade": len(selecionados),
        "total_boletos": round(sum(gasto(r) for o, r, _ in selecionados if o == "boleto"), 2),
        "total_pix": round(sum(gasto(r) for o, r, _ in selecionados if o == "pix"), 2),
        "total_dinheiro": round(sum(gasto(r) for o, r, _ in selecionados if o == "dinheiro"), 2),
        "juros_multa": round(sum(_juros(r) for o, r, _ in selecionados if situacao(r, hoje) == "pago"), 2),
        "pagos_com_juros": sum(1 for o, r, _ in selecionados if situacao(r, hoje) == "pago" and _juros(r)),
        "sem_centro": round(por_centro.get(SEM_CENTRO.lower(), {}).get("total", 0.0), 2),
        "centros": centros,
        "series": principais + ([OUTROS_CENTROS] if len(centros) > 6 else []),
        "serie_mensal": serie_mensal,
        "maiores": maiores,
    }


# --------------------------------------------------------------------------- exportação

_ROTULO_SITUACAO = {
    "vencido": "Vencido", "vence_hoje": "Vence hoje", "a_vencer": "A vencer",
    "pendente": "Pendente", "pago": "Pago", "cancelado": "Cancelado",
}


def _valor_exportacao(chave: str, valor):
    if valor is None:
        return ""
    if chave in ROTULOS:
        return ROTULOS[chave].get(valor, valor)
    return valor


def gerar_csv(registros: list[dict], hoje: date | None = None, colunas=COLUNAS) -> bytes:
    hoje = hoje or date.today()
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    escritor.writerow([c.titulo for c in colunas] + ["Situação"])
    for r in registros:
        linha = []
        for c in colunas:
            v = _valor_exportacao(c.chave, r.get(c.chave))
            if isinstance(v, datetime):
                v = v.strftime("%d/%m/%Y %H:%M")
            elif isinstance(v, date):
                v = v.strftime("%d/%m/%Y")
            elif isinstance(v, float):
                v = f"{v:.2f}".replace(".", ",")
            linha.append(v)
        linha.append(_ROTULO_SITUACAO[situacao(r, hoje)])
        escritor.writerow(linha)
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def _cabecalho(ws, titulos: list[str], larguras: list[int]) -> None:
    for i, (titulo, largura) in enumerate(zip(titulos, larguras), start=1):
        celula = ws.cell(row=1, column=i, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F3D3E")
        celula.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = "A2"


def gerar_xlsx(registros: list[dict], hoje: date | None = None, colunas=COLUNAS, titulo_aba: str = "Boletos") -> bytes:
    hoje = hoje or date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = titulo_aba
    titulos = [c.titulo for c in colunas] + ["Situação"]
    _cabecalho(ws, titulos, [c.largura for c in colunas] + [14])
    for n, r in enumerate(registros, start=2):
        for i, c in enumerate(colunas, start=1):
            celula = ws.cell(row=n, column=i, value=_valor_exportacao(c.chave, r.get(c.chave)) or None)
            if c.formato:
                celula.number_format = c.formato
        ws.cell(row=n, column=len(colunas) + 1, value=_ROTULO_SITUACAO[situacao(r, hoje)])
    total_linha = len(registros) + 2
    col_valor = [c.chave for c in colunas].index("valor") + 1
    letra = get_column_letter(col_valor)
    ws.cell(row=total_linha, column=col_valor - 1, value="Total").font = Font(bold=True)
    total = ws.cell(row=total_linha, column=col_valor, value=f"=SUM({letra}2:{letra}{max(total_linha - 1, 2)})")
    total.font = Font(bold=True)
    total.number_format = '"R$" #,##0.00'
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{max(len(registros) + 1, 1)}"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_xlsx_analise(analise: dict, descricao_filtros: str) -> bytes:
    moeda = '"R$" #,##0.00'
    wb = Workbook()
    ws = wb.active
    ws.title = "Por centro de custo"
    ws.cell(row=1, column=1, value=f"Análise de gastos · {descricao_filtros}").font = Font(bold=True, size=13, color="0F3D3E")
    titulos = ["Centro de custo", "Boletos (R$)", "Pix (R$)", "Total (R$)", "% do total", "Lançamentos", "Pagos (R$)", "Pendentes (R$)"]
    for i, (titulo, largura) in enumerate(zip(titulos, [32, 16, 16, 16, 12, 13, 16, 16]), start=1):
        celula = ws.cell(row=3, column=i, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="0F3D3E")
        ws.column_dimensions[get_column_letter(i)].width = largura
    for n, c in enumerate(analise["centros"], start=4):
        valores = [c["centro"], c["boletos"], c["pix"], c["total"], c["percentual"] / 100, c["quantidade"], c["pagos"], c["pendentes"]]
        for i, valor in enumerate(valores, start=1):
            celula = ws.cell(row=n, column=i, value=valor)
            if i in (2, 3, 4, 7, 8):
                celula.number_format = moeda
            elif i == 5:
                celula.number_format = "0.0%"
    linha_total = len(analise["centros"]) + 4
    ws.cell(row=linha_total, column=1, value="Total").font = Font(bold=True)
    for i in (2, 3, 4, 6, 7, 8):
        letra = get_column_letter(i)
        celula = ws.cell(row=linha_total, column=i, value=f"=SUM({letra}4:{letra}{max(linha_total - 1, 4)})")
        celula.font = Font(bold=True)
        if i != 6:
            celula.number_format = moeda

    ws_mes = wb.create_sheet("Por mês")
    series = analise["series"]
    _cabecalho(ws_mes, ["Mês"] + series + ["Total"], [12] + [22] * len(series) + [16])
    for n, linha in enumerate(analise["serie_mensal"], start=2):
        ws_mes.cell(row=n, column=1, value=linha["rotulo"])
        for i, serie in enumerate(series, start=2):
            ws_mes.cell(row=n, column=i, value=linha["centros"].get(serie, 0)).number_format = moeda
        ws_mes.cell(row=n, column=len(series) + 2, value=linha["total"]).number_format = moeda

    ws_nomes = wb.create_sheet("Maiores beneficiários")
    _cabecalho(ws_nomes, ["Beneficiário / Recebedor", "Total (R$)", "Lançamentos"], [44, 16, 13])
    for n, item in enumerate(analise["maiores"], start=2):
        ws_nomes.cell(row=n, column=1, value=item["nome"])
        ws_nomes.cell(row=n, column=2, value=item["total"]).number_format = moeda
        ws_nomes.cell(row=n, column=3, value=item["quantidade"])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- previsão de saídas


def previsao_saidas(boletos: list[dict], pix: list[dict], hoje: date | None = None, semanas: int = 8) -> dict:
    """Quanto vai sair nas próximas semanas (boletos e Pix pendentes), com o que já está em atraso à parte.

    Cada semana começa hoje (semana 1 = hoje até daqui a 6 dias), para ler "o que vence nos próximos 7 dias".
    """
    hoje = hoje or date.today()
    semanas = max(1, min(int(semanas), 26))
    faixas = [{"inicio": hoje + timedelta(days=7 * n), "fim": hoje + timedelta(days=7 * n + 6)} for n in range(semanas)]
    atraso = {"total": 0.0, "quantidade": 0, "itens": []}
    grupos = [{**f, "total": 0.0, "boletos": 0.0, "pix": 0.0, "quantidade": 0, "itens": []} for f in faixas]
    sem_data = {"total": 0.0, "quantidade": 0}
    for tipo, registros in (("boleto", boletos), ("pix", pix)):
        for r in registros:
            if r.get("status") != "Pendente" or r.get("id") is None or r["id"] < 0:
                continue
            valor = _valor(r)
            venc = r.get("vencimento")
            if not isinstance(venc, date):
                sem_data["total"] += valor
                sem_data["quantidade"] += 1
                continue
            item = {"tipo": tipo, "id": r["id"], "nome": (r.get("recebedor") if tipo == "pix" else r.get("beneficiario")) or "",
                    "valor": round(valor, 2), "vencimento": venc.isoformat(), "centro_custo": r.get("centro_custo") or ""}
            if venc < hoje:
                atraso["total"] += valor
                atraso["quantidade"] += 1
                atraso["itens"].append(item)
                continue
            indice = (venc - hoje).days // 7
            if indice < semanas:
                grupo = grupos[indice]
                grupo["total"] += valor
                grupo["boletos" if tipo == "boleto" else "pix"] += valor
                grupo["quantidade"] += 1
                grupo["itens"].append(item)
    for grupo in grupos:
        grupo["itens"].sort(key=lambda i: (i["vencimento"], -i["valor"]))
        grupo.update(inicio=grupo["inicio"].isoformat(), fim=grupo["fim"].isoformat(),
                     total=round(grupo["total"], 2), boletos=round(grupo["boletos"], 2), pix=round(grupo["pix"], 2))
    atraso["itens"].sort(key=lambda i: i["vencimento"])
    atraso["total"] = round(atraso["total"], 2)
    proximos = lambda dias: round(sum(i["valor"] for g in grupos for i in g["itens"]  # noqa: E731
                                      if date.fromisoformat(i["vencimento"]) < hoje + timedelta(days=dias)), 2)
    return {
        "hoje": hoje.isoformat(),
        "semanas": grupos,
        "em_atraso": atraso,
        "sem_vencimento": {"total": round(sem_data["total"], 2), "quantidade": sem_data["quantidade"]},
        "totais": {"7_dias": proximos(7), "30_dias": proximos(30), "periodo": round(sum(g["total"] for g in grupos), 2),
                   "com_atraso": round(atraso["total"] + sum(g["total"] for g in grupos), 2)},
    }
