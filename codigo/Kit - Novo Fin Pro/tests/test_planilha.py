from datetime import date, datetime

from openpyxl import load_workbook

from app.planilha import PlanilhaBoletos, RegistroNaoEncontrado


def test_ciclo_completo(tmp_path):
    p = PlanilhaBoletos(tmp_path / "boletos.xlsx", tmp_path / "backups")
    a = p.adicionar({"linha_digitavel": "1", "codigo_barras": "ABC", "beneficiario": "Fornecedor A",
                     "vencimento": date(2026, 10, 1), "valor": 99.9, "tipo": "bancario"})
    b = p.adicionar({"beneficiario": "Fornecedor B", "valor": 10.0})
    assert (a["id"], b["id"]) == (1, 2)
    assert isinstance(a["data_lancamento"], datetime)

    registros = p.listar()
    assert [r["beneficiario"] for r in registros] == ["Fornecedor A", "Fornecedor B"]
    assert registros[0]["tipo"] == "bancario" and registros[0]["vencimento"] == date(2026, 10, 1)

    atualizado = p.atualizar(1, {"status": "Pago", "data_pagamento": date(2026, 9, 30)})
    assert atualizado["status"] == "Pago" and atualizado["data_pagamento"] == date(2026, 9, 30)
    assert p.buscar_por_codigo("ABC")[0]["id"] == 1

    p.excluir(2)
    assert [r["id"] for r in p.listar()] == [1]
    assert p.adicionar({"beneficiario": "C"})["id"] == 2

    try:
        p.obter(99)
    except RegistroNaoEncontrado:
        pass
    else:
        raise AssertionError("esperava RegistroNaoEncontrado")

    wb = load_workbook(tmp_path / "boletos.xlsx")
    assert wb.sheetnames == ["Boletos", "Pix", "Caixa", "Resumo", "Centros de Custo", "Histórico"]
    cabecalho = [c.value for c in wb["Boletos"][1]]
    assert "Data de Lançamento" in cabecalho and "Linha Digitável" in cabecalho
    assert (tmp_path / "backups").exists()


def test_lote_em_uma_gravacao_e_aba_historico(tmp_path):
    p = PlanilhaBoletos(tmp_path / "b.xlsx")
    criados = p.adicionar_lote([{"beneficiario": "A", "valor": 1}, {"beneficiario": "B", "valor": 2}])
    assert [c["id"] for c in criados] == [1, 2]

    resultados = p.aplicar_lote([
        {"tipo": "atualizar", "id": 1, "dados": {"status": "Pago", "data_pagamento": date(2026, 9, 14)}},
        {"tipo": "excluir", "id": 2},
        {"tipo": "atualizar", "id": 99, "dados": {"status": "Pago"}},
        {"tipo": "adicionar", "dados": {"beneficiario": "C"}},
    ])
    assert [r["ok"] for r in resultados] == [True, True, False, True]
    assert resultados[3]["registro"]["id"] == 3
    assert [(r["id"], r["beneficiario"], r["status"]) for r in p.listar()] == [(1, "A", "Pago"), (3, "C", "Pendente")]

    historico = list(load_workbook(tmp_path / "b.xlsx")["Histórico"].iter_rows(min_row=2, values_only=True))
    assert [h[1] for h in historico] == ["Lançamento", "Lançamento", "Pagamento", "Exclusão", "Lançamento"]
    assert "Status: Pendente → Pago" in historico[2][6]


def test_ignora_status_desconhecido_e_linhas_vazias(tmp_path):
    caminho = tmp_path / "b.xlsx"
    p = PlanilhaBoletos(caminho)
    p.adicionar({"beneficiario": "X"})
    wb = load_workbook(caminho)
    ws = wb["Boletos"]
    ws.cell(row=2, column=12, value="qualquer")
    ws.cell(row=5, column=5, value=None)
    wb.save(caminho)
    assert p.listar()[0]["status"] == "Pendente"


def test_voltar_para_pendente_apaga_a_data_de_pagamento_gravada(tmp_path):
    """Regressão: apagar um campo precisa limpar a célula (antes a data de pagamento antiga ficava na planilha)."""
    from datetime import date as _date

    planilha = PlanilhaBoletos(tmp_path / "boletos.xlsx")
    registro = planilha.adicionar({"beneficiario": "Fornecedor", "valor": 10, "status": "Pago", "data_pagamento": _date.today()},
                                  usuario="Ana")
    assert planilha.obter(registro["id"])["pago_por"] == "Ana"
    planilha.atualizar(registro["id"], {"status": "Pendente", "data_pagamento": None}, usuario="Ana")
    relido = planilha.obter(registro["id"])
    assert relido["data_pagamento"] is None and relido["pago_por"] == ""
