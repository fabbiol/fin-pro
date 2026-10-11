"""Relatório do diesel em três níveis (simples, detalhado e completo): contas, textos do Telegram, API e Excel."""

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app import contexto
from app.banco import BancoBoletos
from app.diesel import precos, relatorio, servico
from app.diesel.bot import BotDiesel
from app.diesel.modelos import ABAS_DIESEL
from app.notificacoes import canais

SEGUNDA = date(2026, 9, 21)  # semana de 21 a 27/09
QUARTA = date(2026, 9, 23)


@pytest.fixture(autouse=True)
def banco(tmp_path, monkeypatch):
    from app import config

    monkeypatch.setattr(config, "DADOS_DIR", tmp_path)
    diesel = BancoBoletos(tmp_path / "diesel.db", tmp_path / "diesel.xlsx", tmp_path / "backups",
                          abas=ABAS_DIESEL, centros=False, nome_backup="diesel")
    diesel.garantir()
    monkeypatch.setattr(contexto, "diesel", diesel)
    yield diesel
    diesel.parar()


@pytest.fixture
def semana():
    """Uma semana com S10 e comum, histórico de 90 dias antes e um caminhão que gastou muito acima do normal."""
    precos.salvar({"Diesel S10": 7.45, "Diesel comum": 6.75}, "Fábio")
    servico.salvar_maquina({"nome": "Escavadeira", "tipo": "Máquina", "combustivel": "Diesel comum"}, "Fábio")  # horas
    servico.salvar_maquina({"nome": "Caminhão", "tipo": "Caminhão", "combustivel": "Diesel S10"}, "Fábio")  # km
    servico.salvar_maquina({"nome": "Gerador", "tipo": "Gerador", "combustivel": "Diesel comum"}, "Fábio")
    servico.lancar_compra({"data": SEGUNDA, "litros": 1000, "valor": 7450, "combustivel": "Diesel S10"}, "Fábio")
    # histórico: o caminhão já abastecia desde agosto (o "normal" dele é ~4,8 L/dia)
    servico.lancar_abastecimento({"data": date(2026, 8, 10), "maquina": "Caminhão", "litros": 100, "medidor": 140000}, "Almir")
    # semana anterior: escavadeira 150 L, caminhão 100 L
    servico.lancar_abastecimento({"data": date(2026, 9, 15), "maquina": "Escavadeira", "litros": 150, "medidor": 1000}, "Almir")
    servico.lancar_abastecimento({"data": date(2026, 9, 16), "maquina": "Caminhão", "litros": 100, "medidor": 150000}, "Almir")
    # esta semana
    servico.lancar_abastecimento({"data": SEGUNDA, "maquina": "Escavadeira", "litros": 160, "medidor": 1020,
                                  "posto": "Posto City", "pessoa": "Marcos"}, "Almir")
    servico.lancar_abastecimento({"data": QUARTA, "maquina": "Caminhão", "litros": 200, "medidor": 150800,
                                  "posto": "Posto City", "pessoa": "Juliana"}, "Almir")
    servico.lancar_abastecimento({"data": QUARTA, "maquina": "Caminhão", "litros": 100, "posto": "Auto Posto",
                                  "pessoa": "Juliana"}, "Almir")  # sem km


def test_analise_compara_mostra_trechos_e_pontos_de_atencao(semana):
    r = relatorio.analise("semana", SEGUNDA, alerta_valor=1000, hoje=date(2026, 9, 25))
    assert r["abastecido"]["custo"] == 160 * 6.75 + 300 * 7.45
    assert r["anterior"]["custo"] == 150 * 6.75 + 100 * 7.45 and r["variacao"]["litros"] == 84.0  # 460 L contra 250 L
    caminhao = next(m for m in r["por_maquina"] if m["maquina"] == "Caminhão")
    assert caminhao["litros_anterior"] == 100 and caminhao["variacao"] == 200.0
    assert caminhao["media_normal"] == 4.76 and caminhao["acima_do_normal"] > 0  # 60 L/dia contra 4,76 L/dia
    # o trecho: 800 km desde a leitura anterior, com 200 L = 4 km/L; a escavadeira: 20 h com 160 L = 8 L/h
    trechos = {a["maquina"]: a["trecho"] for a in r["abastecimentos"] if a.get("trecho")}
    assert trechos["Caminhão"]["rodado"] == 800 and trechos["Caminhão"]["rotulo"] == "4.0 km/L".replace(".", ",")
    assert trechos["Escavadeira"]["rotulo"] == "8.0 L/h".replace(".", ",")
    textos = " ".join(a["texto"] for a in r["alertas"])
    assert "Caminhão consumiu" in textos and "sem horímetro/km: Caminhão (1)" in textos and "Gerador" in textos
    assert [d["rotulo"] for d in r["serie"]][:3] == ["seg 21", "ter 22", "qua 23"]
    assert r["serie"][2]["por_combustivel"] == {"Diesel S10": 2235.0}
    assert [p["nome"] for p in r["por_posto"]] == ["Posto City", "Auto Posto"]
    assert r["por_pessoa"][0]["nome"] == "Juliana" and r["por_pessoa"][0]["litros"] == 300
    assert r["previsao"]["saldo"] == 7450 - 150 * 6.75 - 200 * 7.45 - 160 * 6.75 - 300 * 7.45
    escavadeira = next(m for m in r["por_maquina"] if m["maquina"] == "Escavadeira")
    assert escavadeira["media_normal"] == 0 and not escavadeira["acima_do_normal"]  # só 6 dias de histórico: sem "normal"
    assert r["previsao"]["dias"] > 0


def test_periodo_passado_nao_tem_previsao(semana):
    r = relatorio.analise("semana", SEGUNDA, hoje=date(2026, 10, 20))
    assert r["previsao"] is None and r["periodo_atual"] is False


def test_tres_niveis_de_texto(semana):
    hoje = date(2026, 9, 25)
    simples = relatorio.texto("semana", SEGUNDA, "simples", 1000, hoje)
    assert "Relatório semanal" in simples and "Abastecido: <b>R$ 3.315,00</b>" in simples
    assert "• Caminhão: 300 L · R$ 2.235,00" in simples and "dá para ~" in simples
    assert "ponto(s) de atenção" in simples and "Pontos de atenção" not in simples
    detalhado = relatorio.texto("semana", SEGUNDA, "detalhado", 1000, hoje)
    assert "Pontos de atenção" in detalhado and "Comparado com" in detalhado and "Por máquina" in detalhado
    completo = relatorio.texto("semana", SEGUNDA, "completo", 1000, hoje)
    assert "<b>Abastecimentos</b>" in completo and "23/09 Caminhão: 200 L" in completo and "4,0 km/L" in completo
    assert "<b>Por posto</b>" in completo and "<b>Quem abasteceu</b>" in completo
    assert len(simples) < len(detalhado) < len(completo)
    assert relatorio.texto("semana", SEGUNDA, "qualquer", 1000, hoje) == simples  # nível desconhecido: simples


def test_texto_longo_e_cortado_para_o_telegram():
    texto = "\n".join(f"linha {n} " + "x" * 60 for n in range(200))
    cortado = relatorio._limitar(texto)
    assert len(cortado) <= relatorio.LIMITE_TELEGRAM and cortado.endswith("(continua no Excel: Diesel → Relatórios)")


@pytest.fixture
def bot(tmp_path, monkeypatch):
    enviados = []

    def telegram(token, metodo, dados=None):
        enviados.append((metodo, dados or {}))
        return {"message_id": 1} if metodo.startswith("send") else True

    monkeypatch.setattr(canais, "_telegram", telegram)
    b = BotDiesel(tmp_path / "diesel_bot.json", lambda dados, nome: {"id": "a1", "nome": nome, "caminho": "arquivos/a1.jpg"},
                  lambda: {"100": "Fábio"}, ler=lambda dados, nome, legenda: {})
    b.configurar({"token": "123:abc", "gestores": ["100"]})
    b.enviados = enviados
    return b


def test_comando_do_bot_e_nivel_dos_automaticos(bot, semana):
    bot._comando("100", "/relatorio mes completo", "Fábio")
    assert "Relatório mensal" in bot.enviados[-1][1]["text"] and "Pontos de atenção" in bot.enviados[-1][1]["text"]
    bot._comando("100", "/relatorio", "Fábio")
    assert "Relatório semanal" in bot.enviados[-1][1]["text"] and "Pontos de atenção" not in bot.enviados[-1][1]["text"]
    assert bot.nivel_relatorio() == "simples"
    bot.configurar({"nivel_relatorio": "detalhado"})
    assert bot.nivel_relatorio() == "detalhado" and bot.config_publica()["nivel_relatorio"] == "detalhado"
    bot.configurar({"nivel_relatorio": "invalido"})
    assert bot.nivel_relatorio() == "detalhado"


def test_api_do_relatorio_e_excel_completo(bot, semana, monkeypatch):
    from fastapi.testclient import TestClient

    from app.main import app

    monkeypatch.setattr(contexto, "diesel_bot", bot)
    with TestClient(app) as cliente:
        r = cliente.get("/api/diesel/relatorio", params={"periodo": "semana", "data": "2026-09-23"}).json()
        assert {"serie", "por_posto", "por_pessoa", "alertas", "anterior", "variacao", "previsao", "precos"} <= set(r)
        assert r["alerta_valor"] == servico.ALERTA_PADRAO
        texto = cliente.get("/api/diesel/relatorio/texto", params={"periodo": "semana", "data": "2026-09-23",
                                                                   "nivel": "completo"}).json()["texto"]
        assert "<b>Abastecimentos</b>" in texto
        excel = cliente.get("/api/diesel/relatorio/arquivo", params={"periodo": "semana", "data": "2026-09-23"})
        wb = load_workbook(io.BytesIO(excel.content))
        assert wb.sheetnames == ["Resumo", "Máquinas", "Por dia", "Combustíveis", "Por posto", "Por pessoa",
                                 "Abastecimentos", "Compras", "Pontos de atenção"]
        assert wb["Resumo"]["B3"].value == 3315
        enviado = cliente.post("/api/diesel/relatorio/enviar", params={"periodo": "semana", "data": "2026-09-23",
                                                                       "nivel": "detalhado"})
        assert enviado.json() == {"enviado": 1} and "Pontos de atenção" in bot.enviados[-1][1]["text"]
